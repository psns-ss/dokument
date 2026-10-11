# logic/parser.py
# Reads the table in descr.docx and returns a dict of variable descriptions.
# Table cells:
# 0 - variable name, 1 - description, 2 - show in form (❌ = no),
# 3 - type (text/enum/bool/match/subdoc/table, comma-separated),
# 4 - options: typed form <type>:<value>[;<type>:<value>]*
#     (for any number of types) or a bare value - single-type only.
# Value per type: enum - dot-path to a list in language_rules.yaml,
# match - YAML filename for the match plugin, text - case codes,
# file - a number (max uploads, 0 = no cap).
# Result: var_meta["options"] = {type: value} (dict; None = no options).

import re
import yaml
from pathlib import Path
from typing import Dict, Any
from docx import Document
from docx.text.paragraph import Paragraph
import logging

from logic.inflection import RECOGNIZED_CASES

logger = logging.getLogger(__name__)

# Sentinel description for packs whose descr.docx has no leading paragraph.
# Shared with pack_manager, which must NOT hash it — two undescribed packs would
# otherwise collide on the same hash and be falsely flagged as duplicates.
NO_DESCRIPTION = "No description available"

# Types that read the options column (each has its own value format).
_CONSUMER_TYPES = ("enum", "match", "text", "file")

# All known types — for recognizing the prefix in the typed form.
_KNOWN_TYPES = {"text", "enum", "bool", "match", "computed",
                "subdoc", "form", "file", "table"}

# A typed-form segment <type>:<value> (value non-empty).
_TYPED_PAIR = re.compile(r"^([a-z]+):(.+)$", re.DOTALL)


def _load_lang_rules(lang_file: Path) -> Dict[str, Any]:
    """Load language_rules.yaml, if it exists. Degrade-and-log (never raise):
    a bad file must not 500 the form — enum options then just degrade to []."""
    if not lang_file.exists():
        return {}
    try:
        with lang_file.open(encoding="utf-8") as f:
            return yaml.safe_load(f) or {}
    except (yaml.YAMLError, OSError) as e:
        logger.error(f"language_rules: failed to load {lang_file}: {e}; using empty rules")
        return {}


def _parse_options_cell(var: str, types: list, cell: str,
                        lang_rules: Dict[str, Any]) -> Dict[str, Any]:
    """Parse the options-column cell → {type: value} (empty dict = none).

    Typed form <type>:<value>[;...] — for any number of types; a bare value —
    single-type only (multi-type is an error: a type may later gain its own
    options, and the bare cell's meaning would change).
    Per-segment: if at least one segment is a valid 'type:value' pair the cell
    is typed — good segments are kept, malformed ones are logged + skipped (not
    all-or-nothing). Errors — log + degrade, no raise: the parser is on the
    request path, a descr typo must not 500 the form.
    """
    cell = cell.strip()
    if not cell or cell == "-":
        return {}

    pairs = []
    bad = []
    for segment in cell.split(";"):
        seg = segment.strip()
        if not seg:
            continue
        m = _TYPED_PAIR.match(seg)
        if m and m.group(1) in _KNOWN_TYPES:
            pairs.append((m.group(1), m.group(2).strip()))
        else:
            bad.append(seg)

    if pairs:
        # Typed cell — keep the good segments, log + skip the malformed ones.
        for seg in bad:
            logger.error(
                f"descr options: var '{var}' — segment {seg!r} is not a valid "
                f"'type:value' pair; skipped (the other segments are kept)"
            )
        return _apply_pairs(var, types, pairs, lang_rules)

    # No typed segment survived — the cell is bare (single-type OK; multi-type
    # is an error in _bare_cell).
    return _bare_cell(var, types, cell, lang_rules)


def _bare_cell(var: str, types: list, cell: str,
               lang_rules: Dict[str, Any]) -> Dict[str, Any]:
    """Bare value: single-type — as before; multi-type — error."""
    if len(types) > 1:
        logger.error(
            f"descr options: var '{var}' — bare cell on multi-type var "
            f"({', '.join(types)}), use 'type:value' pairs; cell: {cell!r}"
        )
        return {}
    return _apply_pairs(var, types, [(types[0], cell)], lang_rules)


def _apply_pairs(var: str, types: list, pairs: list,
                 lang_rules: Dict[str, Any]) -> Dict[str, Any]:
    """Each <type>:<value> pair goes to its type's parser."""
    options: Dict[str, Any] = {}
    for t, value in pairs:
        if t not in types:
            logger.error(
                f"descr options: var '{var}' — options for type '{t}' not in "
                f"the var's types ({', '.join(types)}); pair skipped; "
                f"value: {value!r}"
            )
            continue
        if t not in _CONSUMER_TYPES:
            logger.debug(
                f"descr options: var '{var}' — type '{t}' has no options; "
                f"pair skipped"
            )
            continue
        if t in options:
            logger.warning(
                f"descr options: var '{var}' — duplicate '{t}' pair; last wins"
            )
        options[t] = _parse_value(var, t, value, lang_rules)
    return options


def _parse_value(var: str, t: str, value: str,
                 lang_rules: Dict[str, Any]) -> Any:
    if t == "enum":
        return _parse_enum_options(var, value, lang_rules)
    if t == "text":
        return _parse_text_cases(var, value)
    if t == "file":
        return _parse_file_limit(var, value)
    return value  # match: the key as-is (non-empty — empty is filtered above)


def _parse_enum_options(var: str, path: str,
                        lang_rules: Dict[str, Any]) -> list:
    """Dot-path in language_rules.yaml → list of options."""
    try:
        value = lang_rules
        for part in path.split("."):
            value = value[part]
        if isinstance(value, list):
            return value
        logger.error(
            f"descr options: var '{var}' — path {path!r} is not a list in "
            f"language_rules.yaml"
        )
    except (KeyError, TypeError):
        logger.error(
            f"descr options: var '{var}' — path {path!r} not found in "
            f"language_rules.yaml"
        )
    return []


def _parse_text_cases(var: str, value: str) -> list:
    """CSV of case codes; unknown codes are dropped (with a log)."""
    codes = [c.strip() for c in value.split(",") if c.strip()]
    bad = [c for c in codes if c not in RECOGNIZED_CASES]
    if bad:
        logger.error(
            f"descr options: var '{var}' — unknown case code(s) {bad}; "
            f"cell: {value!r}"
        )
    return [c for c in codes if c in RECOGNIZED_CASES]


def _parse_file_limit(var: str, value: str) -> int:
    """file:<N> — max uploads (positive int). 0 = no cap
    (degrade on error: the parser is on the request path, don't 500 the form)."""
    try:
        n = int(value)
    except (TypeError, ValueError):
        logger.error(
            f"descr options: var '{var}' — file limit {value!r} is not an int; "
            f"no cap"
        )
        return 0
    if n <= 0:
        logger.error(
            f"descr options: var '{var}' — file limit {value!r} must be positive; "
            f"no cap"
        )
        return 0
    return n


def read_variable_descriptions(descr_path: Path, lang_file: Path) -> Dict[str, Dict[str, Any]]:
    """
    Parse the DOCX table and return a dict:
    { "var_name": {"description": "...", "show_in_form": True|False,
                   "type": "text|enum", "options": {type: value}}, ... }
    """
    lang_rules = _load_lang_rules(lang_file)
    doc = Document(descr_path)
    variables: Dict[str, Dict[str, Any]] = {}

    # Iterate over all tables
    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            if len(cells) < 4:
                continue  # skip incomplete rows

            # Variable name without braces
            var = cells[0].text.strip().replace("{", "").replace("}", "").strip()
            if not var:
                continue

            # Fill in the metadata
            type_raw = cells[3].text.strip().lower() or "text"
            # ❌ = do not show; strip the U+FE0F variation selector so the
            # emoji-presentation form (❌︎) is recognized too.
            show_cell = cells[2].text.strip().replace(chr(0xFE0F), "")
            var_meta = {
                "description": cells[1].text.strip(),
                "show_in_form": show_cell != "❌",  # ❌ = do not show
                "type": type_raw,
                "types": [t.strip() for t in type_raw.split(',')] if ',' in type_raw else [type_raw],
                "options": None,
            }

            # Options column: typed form or bare value (single-type)
            if len(cells) > 4:
                options = _parse_options_cell(var, var_meta["types"],
                                              cells[4].text, lang_rules)
                if options:
                    var_meta["options"] = options

            variables[var] = var_meta
    return variables


def extract_pack_description(descr_path: Path) -> str:
    """
    Extract the pack description from the first paragraph before the table.

    Returns the paragraph text, or 'No description available' if none exists.
    """
    doc = Document(descr_path)

    # Find first paragraph that comes before any table
    for element in doc.element.body:
        if element.tag.endswith('tbl'):  # Table element - stop at first table
            break
        if element.tag.endswith('p'):  # Paragraph element
            para = Paragraph(element, doc)  # Bind to original document
            text = para.text.strip()
            if text:
                return text

    return NO_DESCRIPTION
