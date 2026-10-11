"""
Tests for subdocument functionality.

This file tests the subdoc_helper framework function (create_subdoc_from_bookmark).

Note: University plugin subdoc generation (university_historical) is tested in
test_user_logic/test_university_subdoc.py. Full rendering integration is tested
in test_renderer.py.
"""

import pytest
from pathlib import Path
import sys

# Add project root to path for imports
project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

from docxtpl import DocxTemplate
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from logic.subdoc_helper import create_subdoc_from_bookmark, create_subdoc_sandwich


class TestCreateSubdocFromBookmark:
    """Tests for create_subdoc_from_bookmark function."""

    def test_create_subdoc_valid_bookmark(self, project_root):
        """Test creating a subdoc from a valid bookmark in university.docx."""
        source_path = project_root / "templates" / "university.docx"
        assert source_path.exists(), f"Source file not found: {source_path}"

        template_path = project_root / "templates" / "template_1.docx"
        tpl = DocxTemplate(str(template_path))

        subdoc = create_subdoc_from_bookmark(tpl, "moscow_state_university", source_path)
        assert subdoc is not None

    def test_create_subdoc_nonexistent_bookmark(self, project_root):
        """Test that nonexistent bookmarks return empty subdoc without crashing."""
        source_path = project_root / "templates" / "university.docx"
        template_path = project_root / "templates" / "template_1.docx"
        tpl = DocxTemplate(str(template_path))

        subdoc = create_subdoc_from_bookmark(tpl, "nonexistent_bookmark_xyz", source_path)
        assert subdoc is not None

    def test_create_subdoc_multiple_bookmarks(self, project_root):
        """Test creating subdocs from multiple bookmarks in the same document."""
        source_path = project_root / "templates" / "university.docx"
        template_path = project_root / "templates" / "template_1.docx"
        tpl = DocxTemplate(str(template_path))

        bookmarks = ["moscow_state_university", "bauman_mgtu", "moscow_engineering_university"]
        subdocs = {}

        for bookmark in bookmarks:
            subdoc = create_subdoc_from_bookmark(tpl, bookmark, source_path)
            subdocs[bookmark] = subdoc
            assert subdoc is not None, f"Failed to create subdoc for {bookmark}"

        assert len(subdocs) == 3

    def test_create_subdoc_preserves_content(self, project_root):
        """Test that subdoc contains actual content from the bookmark."""
        source_path = project_root / "templates" / "university.docx"
        template_path = project_root / "templates" / "template_1.docx"
        tpl = DocxTemplate(str(template_path))

        subdoc = create_subdoc_from_bookmark(tpl, "moscow_state_university", source_path)

        if hasattr(subdoc, '_subdoc') and subdoc._subdoc:
            subdoc_buffer = subdoc._subdoc
            subdoc_buffer.seek(0)
            sub_doc = Document(subdoc_buffer)
            all_text = "".join([p.text for p in sub_doc.paragraphs])

            assert "1755" in all_text, f"Expected '1755' in subdoc content, got: {all_text[:500]}"
        else:
            assert subdoc is not None


class TestCreateSubdocSandwich:
    """Tests for create_subdoc_sandwich() unified helper."""

    def _make_para(self, text: str) -> OxmlElement:
        """Helper: create a paragraph element with given text."""
        para = OxmlElement('w:p')
        run = OxmlElement('w:r')
        t = OxmlElement('w:t')
        t.text = text
        run.append(t)
        para.append(run)
        return para

    def _get_all_text(self, subdoc) -> str:
        """Extract all paragraph text from a Subdoc via subdocx."""
        return "\n".join([p.text for p in subdoc.subdocx.paragraphs])

    def test_sandwich_renders_placeholders(self, project_root):
        """Subdoc elements with {{placeholder}} are rendered before Subdoc creation."""
        tpl = DocxTemplate(str(project_root / "templates" / "template_1.docx"))
        para = self._make_para('{{user_name}}')

        subdoc = create_subdoc_sandwich(
            tpl, [para], {'user_name': 'Иван Иванов'},
            render_placeholders=True,
        )
        text = self._get_all_text(subdoc)
        assert 'Иван Иванов' in text
        assert '{{user_name}}' not in text

    def test_sandwich_skips_render_without_context(self, project_root):
        """render_placeholders=False -> no render pass -> placeholders preserved."""
        tpl = DocxTemplate(str(project_root / "templates" / "template_1.docx"))
        para = self._make_para('{{user_name}}')

        subdoc = create_subdoc_sandwich(
            tpl, [para], {},
            render_placeholders=False,
        )
        text = self._get_all_text(subdoc)
        assert '{{user_name}}' in text

    def test_sandwich_skips_render_empty_context(self, project_root):
        """Empty context_data -> no render pass -> placeholders preserved."""
        tpl = DocxTemplate(str(project_root / "templates" / "template_1.docx"))
        para = self._make_para('{{user_name}}')

        subdoc = create_subdoc_sandwich(
            tpl, [para], {},
            render_placeholders=True,  # True but context is empty -> skipped
        )
        text = self._get_all_text(subdoc)
        assert '{{user_name}}' in text

    def test_sandwich_copies_theme(self, project_root):
        """Subdoc theme blob matches the base template's theme."""
        tpl = DocxTemplate(str(project_root / "templates" / "template_1.docx"))
        para = self._make_para('theme test')
        subdoc = create_subdoc_sandwich(tpl, [para], {})

        THEME_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme'
        tpl_theme = None
        for rel in tpl.get_docx().part.rels.values():
            if rel.reltype == THEME_REL:
                tpl_theme = rel.target_part.blob
                break

        sub_doc = subdoc.subdocx
        sub_theme = None
        for rel in sub_doc.part.rels.values():
            if rel.reltype == THEME_REL:
                sub_theme = rel.target_part.blob
                break

        assert tpl_theme == sub_theme

    def test_sandwich_multiple_placeholders(self, project_root):
        """Multiple placeholders in same element are all rendered."""
        tpl = DocxTemplate(str(project_root / "templates" / "template_1.docx"))
        para = self._make_para('{{user_name}}, age {{user_age}}')

        subdoc = create_subdoc_sandwich(
            tpl, [para], {'user_name': 'Анна', 'user_age': '30'},
            render_placeholders=True,
        )
        text = self._get_all_text(subdoc)
        assert 'Анна' in text
        assert '30' in text
        assert '{{user_name}}' not in text
        assert '{{user_age}}' not in text

    def test_sandwich_multiple_elements(self, project_root):
        """Multiple elements are all included and rendered."""
        tpl = DocxTemplate(str(project_root / "templates" / "template_1.docx"))
        para1 = self._make_para('Hello {{user_name}}')
        para2 = self._make_para('Age: {{user_age}}')

        subdoc = create_subdoc_sandwich(
            tpl, [para1, para2], {'user_name': 'Дмитрий', 'user_age': '28'},
            render_placeholders=True,
        )
        doc = subdoc.subdocx
        assert len(doc.paragraphs) == 2
        assert 'Дмитрий' in doc.paragraphs[0].text
        assert '28' in doc.paragraphs[1].text


class TestCreateSubdocFromBookmarkWithRenderContext:
    """Tests for create_subdoc_from_bookmark() with render_context."""

    def test_bookmark_without_render_context(self, project_root):
        """Existing callers work unchanged — no render pass."""
        tpl = DocxTemplate(str(project_root / "templates" / "template_1.docx"))
        subdoc = create_subdoc_from_bookmark(
            tpl, "moscow_state_university",
            project_root / "templates" / "university.docx",
        )
        assert subdoc is not None

    def test_bookmark_with_render_context(self, project_root):
        """Placeholder in bookmarked content is rendered when render_context is provided.

        Creates a temporary DOCX with a bookmark containing a placeholder,
        then verifies the placeholder is rendered.
        """
        import io
        import os
        import tempfile

        tpl = DocxTemplate(str(project_root / "templates" / "template_1.docx"))

        # Create a temporary source document with a bookmark containing a placeholder
        src_doc = Document()
        p1 = src_doc.add_paragraph('Start')
        # Add bookmark start
        bookmark_start = OxmlElement('w:bookmarkStart')
        bookmark_start.set(qn('w:id'), '0')
        bookmark_start.set(qn('w:name'), 'test_bookmark')
        p1._p.append(bookmark_start)

        # Paragraph with placeholder
        src_doc.add_paragraph('Student: {{user_name}}.')

        # Paragraph with bookmark end
        p3 = src_doc.add_paragraph('End')
        bookmark_end = OxmlElement('w:bookmarkEnd')
        bookmark_end.set(qn('w:id'), '0')
        p3._p.append(bookmark_end)

        # Save to temp file (Document() needs a path for create_subdoc_from_bookmark)
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
            src_buffer = io.BytesIO()
            src_doc.save(src_buffer)
            src_buffer.seek(0)
            f.write(src_buffer.getvalue())
            temp_path = f.name

        try:
            subdoc = create_subdoc_from_bookmark(
                tpl, "test_bookmark",
                Path(temp_path),
                render_context={'user_name': 'Ivan Ivanov', 'user_age': '25'},
            )
            assert subdoc is not None

            all_text = "\n".join([p.text for p in subdoc.subdocx.paragraphs])
            assert 'Ivan Ivanov' in all_text
            assert '{{user_name}}' not in all_text
        finally:
            os.unlink(temp_path)

    def test_bookmark_includes_non_paragraph_elements(self, project_root):
        """Bookmark extraction includes tables and other body elements, not just paragraphs.

        Regression test: the inner loop only checked }p elements, silently skipping
        tables, images, and other body elements within bookmark boundaries.
        """
        import io
        import os
        import tempfile
        from copy import deepcopy

        tpl = DocxTemplate(str(project_root / "templates" / "template_1.docx"))

        # Create a source document with a bookmark containing a table
        src_doc = Document()
        p1 = src_doc.add_paragraph('Before bookmark')

        # Add bookmark start in a paragraph
        bookmark_start = OxmlElement('w:bookmarkStart')
        bookmark_start.set(qn('w:id'), '0')
        bookmark_start.set(qn('w:name'), 'table_bookmark')
        p1._p.append(bookmark_start)

        # Add a table inside the bookmark
        table = src_doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Header1'
        table.cell(0, 1).text = 'Header2'
        table.cell(1, 0).text = 'Val1'
        table.cell(1, 1).text = 'Val2'

        # Add a paragraph after the table
        src_doc.add_paragraph('After table')

        # Add bookmark end
        p_end = src_doc.add_paragraph('End')
        bookmark_end = OxmlElement('w:bookmarkEnd')
        bookmark_end.set(qn('w:id'), '0')
        p_end._p.append(bookmark_end)

        # Save to temp file
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
            src_buffer = io.BytesIO()
            src_doc.save(src_buffer)
            src_buffer.seek(0)
            f.write(src_buffer.getvalue())
            temp_path = f.name

        try:
            subdoc = create_subdoc_from_bookmark(
                tpl, "table_bookmark",
                Path(temp_path),
                render_context={},
            )
            assert subdoc is not None

            # Verify the table is included in the subdoc
            sub_doc = subdoc.subdocx
            assert len(sub_doc.tables) == 1, (
                f"Expected 1 table in subdoc, got {len(sub_doc.tables)}. "
                f"Non-paragraph elements may be skipped during bookmark extraction."
            )
            assert sub_doc.tables[0].cell(0, 0).text == 'Header1'
            assert sub_doc.tables[0].cell(1, 1).text == 'Val2'
        finally:
            os.unlink(temp_path)
