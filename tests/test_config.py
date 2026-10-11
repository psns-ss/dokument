"""
Tests for app/config.py - configuration with multi-pack support.

Tests dynamic config properties, active pack management, and helpers.
"""

import pytest
from pathlib import Path
from docx import Document

from config import Config, cfg, ensure_active_pack, get_data_context


class TestConfigProperties:
    """Tests for Config class dynamic properties."""

    def test_config_is_not_frozen(self):
        """Config should allow setting _active_pack."""
        # This should not raise - Config is no longer frozen
        cfg.active_pack = "test_pack"
        assert cfg.active_pack == "test_pack"

    def test_active_pack_defaults_to_root(self):
        """active_pack should default to 'root'."""
        # Create a fresh Config
        config = Config()
        assert config.active_pack == "root"

    def test_descr_file_for_root_pack(self):
        """descr_file should resolve to root descr.docx."""
        cfg.active_pack = "root"
        assert cfg.descr_file == cfg.templates_dir / "descr.docx"

    def test_descr_file_for_named_pack(self):
        """descr_file should resolve to pack's descr.docx."""
        cfg.active_pack = "my_pack"
        assert cfg.descr_file == cfg.templates_dir / "my_pack" / "descr.docx"

    def test_pack_dir_for_root_pack(self):
        """pack_dir should be templates_dir for root."""
        cfg.active_pack = "root"
        assert cfg.pack_dir == cfg.templates_dir

    def test_pack_dir_for_named_pack(self):
        """pack_dir should be pack's directory."""
        cfg.active_pack = "my_pack"
        assert cfg.pack_dir == cfg.templates_dir / "my_pack"

    def test_logic_dir_root_uses_root_logic(self, tmp_path):
        """logic_dir delegates to resolve_pack_dirs: root -> root logic/."""
        (tmp_path / "logic").mkdir()
        config = Config(templates_dir=tmp_path)
        config.active_pack = "root"
        assert config.logic_dir == tmp_path / "logic"

    def test_logic_dir_named_pack_falls_back_to_root(self, tmp_path):
        """A named pack with no logic/ of its own falls back to root logic/
        (the effective dir the routes load via resolve_pack_dirs)."""
        (tmp_path / "logic").mkdir()
        config = Config(templates_dir=tmp_path)
        config.active_pack = "my_pack"  # no my_pack/logic on disk
        assert config.logic_dir == tmp_path / "logic"

    def test_logic_dir_named_pack_prefers_own_logic(self, tmp_path):
        """A named pack with its own logic/ uses it (no fallback)."""
        (tmp_path / "logic").mkdir()
        (tmp_path / "my_pack" / "logic").mkdir(parents=True)
        config = Config(templates_dir=tmp_path)
        config.active_pack = "my_pack"
        assert config.logic_dir == tmp_path / "my_pack" / "logic"

    def test_lang_rules_fallback_to_root(self, tmp_path):
        """lang_rules should fallback to root if pack has none."""
        # Create pack without language_rules.yaml
        pack_dir = tmp_path / "my_pack"
        pack_dir.mkdir()

        config = Config(templates_dir=tmp_path)
        config.active_pack = "my_pack"

        # Should fallback to root's language_rules.yaml (which also doesn't exist)
        # but the path should point to root's file
        assert config.lang_rules == tmp_path / "language_rules.yaml"

    def test_base_defaults(self, monkeypatch):
        """Base Config() defaults that are not pack-derived."""
        monkeypatch.delenv('DOCUMENT_STORE', raising=False)
        config = Config()
        assert config.templates_dir == Path("templates")
        assert config.document_store == Path("/document_store")

    def test_document_store_env_override(self, monkeypatch, tmp_path):
        """DOCUMENT_STORE relocates the store — host dev / CI without Docker."""
        target = tmp_path / "store"
        monkeypatch.setenv('DOCUMENT_STORE', str(target))
        assert Config().document_store == target


class TestEnsureActivePack:
    """Tests for ensure_active_pack() helper function."""

    def test_uses_session_active_pack(self):
        """ensure_active_pack should use session's active_pack."""
        session = {"active_pack": "session_pack"}
        result = ensure_active_pack(session)

        assert result == "session_pack"
        assert cfg.active_pack == "session_pack"
        assert session["active_pack"] == "session_pack"

    def test_defaults_to_root_if_descr_exists(self, tmp_path):
        """ensure_active_pack should default to root if descr.docx exists."""
        # Create a descr.docx at root
        descr = tmp_path / "descr.docx"
        doc = Document()
        doc.save(descr)

        session = {}
        config = Config(templates_dir=tmp_path)

        # Temporarily replace cfg for this test
        import config as config_module
        original_cfg = config_module.cfg
        config_module.cfg = config

        try:
            result = ensure_active_pack(session)
            assert result == "root"
            assert config.active_pack == "root"
            assert session["active_pack"] == "root"
        finally:
            config_module.cfg = original_cfg

    def test_returns_none_if_no_descr_at_root(self, tmp_path):
        """ensure_active_pack should return None if no descr.docx at root."""
        session = {}
        config = Config(templates_dir=tmp_path)

        # No descr.docx created

        import config as config_module
        original_cfg = config_module.cfg
        config_module.cfg = config

        try:
            result = ensure_active_pack(session)
            assert result is None
        finally:
            config_module.cfg = original_cfg


class TestGetDataContext:
    """Tests for get_data_context() helper function."""

    def test_raises_error_when_not_initialized(self):
        """get_data_context should raise RuntimeError when not initialized."""
        from logic.context_manager import invalidate_data_context
        invalidate_data_context()

        with pytest.raises(RuntimeError, match="DataContext not initialized"):
            get_data_context()

    def test_returns_loaded_context(self, tmp_path):
        """get_data_context should return the loaded DataContext class."""
        from logic.context_manager import load_data_context_for_pack

        # Create a descr.docx
        doc = Document()
        table = doc.add_table(rows=1, cols=4)
        table.rows[0].cells[0].text = "test_var"
        table.rows[0].cells[1].text = "Test"
        table.rows[0].cells[2].text = "✓"
        table.rows[0].cells[3].text = "text"

        descr_path = tmp_path / "descr.docx"
        doc.save(descr_path)

        lang_path = tmp_path / "language_rules.yaml"
        lang_path.write_text("")

        # Load
        load_data_context_for_pack(descr_path, lang_path)

        # Retrieve via helper
        DataContext = get_data_context()

        assert hasattr(DataContext, "test_var")


class TestComputeOnError:
    """COMPUTE_ON_ERROR env parsing (lenient|strict, default lenient)."""

    def _parse(self):
        from config import _compute_on_error
        return _compute_on_error()

    def test_default_lenient(self, monkeypatch):
        monkeypatch.delenv('COMPUTE_ON_ERROR', raising=False)
        assert self._parse() == 'lenient'

    def test_strict(self, monkeypatch):
        monkeypatch.setenv('COMPUTE_ON_ERROR', 'strict')
        assert self._parse() == 'strict'

    def test_case_insensitive(self, monkeypatch):
        monkeypatch.setenv('COMPUTE_ON_ERROR', 'STRICT')
        assert self._parse() == 'strict'

    def test_unknown_value_falls_back_to_lenient(self, monkeypatch):
        monkeypatch.setenv('COMPUTE_ON_ERROR', 'yell')
        assert self._parse() == 'lenient'
