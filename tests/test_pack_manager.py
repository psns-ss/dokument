"""
Tests for app/pack_manager.py - pack discovery and validation.
"""

import pytest
from pathlib import Path
from docx import Document

from pack_manager import discover_packs, validate_pack, resolve_pack_dirs


def create_descr_with_description(path: Path, description: str, template_count: int = 2):
    """Helper to create a descr.docx with a description paragraph."""
    doc = Document()
    if description:
        doc.add_paragraph(description)

    # Add a table
    table = doc.add_table(rows=2, cols=4)
    table.rows[0].cells[0].text = "var1"
    table.rows[0].cells[1].text = "Variable 1"
    table.rows[0].cells[2].text = "✓"
    table.rows[0].cells[3].text = "text"
    table.rows[1].cells[0].text = "var2"
    table.rows[1].cells[1].text = "Variable 2"
    table.rows[1].cells[2].text = "✓"
    table.rows[1].cells[3].text = "text"

    doc.save(path)

    # Create template files
    for i in range(template_count):
        tpl_path = path.parent / f"template{i+1}.docx"
        tpl_doc = Document()
        tpl_doc.save(tpl_path)


class TestDiscoverPacks:
    """Tests for discover_packs() - discovers all template packs."""

    def test_discover_root_pack_only(self, tmp_path):
        """Should discover root pack when no subdirectories exist."""
        # Create root descr.docx
        create_descr_with_description(tmp_path / "descr.docx", "Root pack description")

        packs = discover_packs(tmp_path)

        assert len(packs) == 1
        assert packs[0]['identifier'] == 'root'
        assert packs[0]['description'] == "Root pack description"
        assert packs[0]['is_disabled'] is False
        assert packs[0]['error'] is None

    def test_discover_root_and_subdir_packs(self, tmp_path):
        """Should discover both root and subdirectory packs."""
        # Create root pack
        create_descr_with_description(tmp_path / "descr.docx", "Root pack")

        # Create subdirectory pack
        subdir = tmp_path / "my_pack"
        subdir.mkdir()
        create_descr_with_description(subdir / "descr.docx", "My pack description")

        packs = discover_packs(tmp_path)

        assert len(packs) == 2
        assert packs[0]['identifier'] == 'root'
        assert packs[1]['identifier'] == 'my_pack'

    def test_root_pack_always_first(self, tmp_path):
        """Root pack should always appear first in the list."""
        # Create multiple subdirectory packs
        for name in ['zebra', 'alpha', 'beta']:
            subdir = tmp_path / name
            subdir.mkdir()
            create_descr_with_description(subdir / "descr.docx", f"{name} pack")

        # Create root pack
        create_descr_with_description(tmp_path / "descr.docx", "Root")

        packs = discover_packs(tmp_path)

        assert packs[0]['identifier'] == 'root'
        # Subdirectories should be sorted alphabetically
        assert packs[1]['identifier'] == 'alpha'
        assert packs[2]['identifier'] == 'beta'
        assert packs[3]['identifier'] == 'zebra'

    def test_skips_logic_directory(self, tmp_path):
        """Should skip the logic/ directory."""
        # Create root pack
        create_descr_with_description(tmp_path / "descr.docx", "Root")

        # Create logic directory (should be skipped)
        logic_dir = tmp_path / "logic"
        logic_dir.mkdir()
        create_descr_with_description(logic_dir / "descr.docx", "Logic pack")

        packs = discover_packs(tmp_path)

        assert len(packs) == 1
        assert packs[0]['identifier'] == 'root'

    def test_skips_pycache_directory(self, tmp_path):
        """Should skip __pycache__ directory."""
        create_descr_with_description(tmp_path / "descr.docx", "Root")

        pycache = tmp_path / "__pycache__"
        pycache.mkdir()
        create_descr_with_description(pycache / "descr.docx", "Cache pack")

        packs = discover_packs(tmp_path)

        assert len(packs) == 1

    def test_skips_subdir_without_descr(self, tmp_path):
        """Should skip subdirectories without descr.docx."""
        create_descr_with_description(tmp_path / "descr.docx", "Root")

        # Create subdir without descr.docx
        subdir = tmp_path / "empty_pack"
        subdir.mkdir()
        # Don't create descr.docx

        packs = discover_packs(tmp_path)

        assert len(packs) == 1
        assert packs[0]['identifier'] == 'root'

    def test_template_count_for_root(self, tmp_path):
        """Should count template*.docx files in root."""
        # Create root with 3 templates
        create_descr_with_description(tmp_path / "descr.docx", "Root", template_count=3)

        packs = discover_packs(tmp_path)

        assert packs[0]['template_count'] == 3

    def test_template_count_for_subdir(self, tmp_path):
        """Should count template*.docx files in subdirectory."""
        create_descr_with_description(tmp_path / "descr.docx", "Root", template_count=1)

        subdir = tmp_path / "my_pack"
        subdir.mkdir()
        create_descr_with_description(subdir / "descr.docx", "My pack", template_count=5)

        packs = discover_packs(tmp_path)

        assert packs[1]['template_count'] == 5

    def test_marks_duplicate_descriptions(self, tmp_path):
        """Should mark packs with duplicate descriptions."""
        # Create root with description "Same"
        create_descr_with_description(tmp_path / "descr.docx", "Same description")

        # Create subdir with same description
        subdir = tmp_path / "my_pack"
        subdir.mkdir()
        create_descr_with_description(subdir / "descr.docx", "Same description")

        packs = discover_packs(tmp_path)

        assert packs[0]['is_duplicate'] is True
        assert packs[1]['is_duplicate'] is True

    def test_no_duplicate_mark_for_unique_descriptions(self, tmp_path):
        """Should not mark packs with unique descriptions."""
        create_descr_with_description(tmp_path / "descr.docx", "Root description")

        subdir = tmp_path / "my_pack"
        subdir.mkdir()
        create_descr_with_description(subdir / "descr.docx", "Different description")

        packs = discover_packs(tmp_path)

        assert packs[0]['is_duplicate'] is False
        assert packs[1]['is_duplicate'] is False

    def test_empty_description_returns_available_message(self, tmp_path):
        """Should return 'No description available' for empty descriptions."""
        # Create descr without description paragraph
        doc = Document()
        table = doc.add_table(rows=1, cols=4)
        table.rows[0].cells[0].text = "var1"
        table.rows[0].cells[1].text = "Var"
        table.rows[0].cells[2].text = "✓"
        table.rows[0].cells[3].text = "text"
        doc.save(tmp_path / "descr.docx")

        packs = discover_packs(tmp_path)

        assert packs[0]['description'] == "No description available"

    def test_returns_empty_list_when_no_root(self, tmp_path):
        """Should return empty list when no root descr.docx exists."""
        packs = discover_packs(tmp_path)

        assert packs == []


class TestValidatePack:
    """Tests for validate_pack() - validates a specific pack."""

    def test_validates_root_pack(self, tmp_path):
        """Should validate root pack successfully."""
        create_descr_with_description(tmp_path / "descr.docx", "Root")

        valid, error = validate_pack(tmp_path, 'root')

        assert valid is True
        assert error is None

    def test_validates_subdir_pack(self, tmp_path):
        """Should validate subdirectory pack successfully."""
        subdir = tmp_path / "my_pack"
        subdir.mkdir()
        create_descr_with_description(subdir / "descr.docx", "My pack")

        valid, error = validate_pack(tmp_path, 'my_pack')

        assert valid is True
        assert error is None

    def test_invalid_pack_not_found(self, tmp_path):
        """Should return error for non-existent pack."""
        create_descr_with_description(tmp_path / "descr.docx", "Root")

        valid, error = validate_pack(tmp_path, 'nonexistent')

        assert valid is False
        assert "not found" in error.lower()

    def test_invalid_root_missing_descr(self, tmp_path):
        """Should return error for root pack without descr.docx."""
        # Don't create descr.docx

        valid, error = validate_pack(tmp_path, 'root')

        assert valid is False
        assert "not found" in error.lower()


class TestResolvePackDirs:
    """Tests for resolve_pack_dirs() - thread-safe (pack_dir, logic_dir) resolution.

    Pure path resolution (no descr.docx needed); logic_dir falls back to the
    root logic/ when the pack has none of its own.
    """

    def test_root_pack_dirs(self, tmp_path):
        """Root: pack_dir = templates_dir, logic_dir = templates_dir/logic."""
        (tmp_path / "logic").mkdir()
        pack_dir, logic_dir = resolve_pack_dirs(tmp_path, 'root')
        assert pack_dir == tmp_path
        assert logic_dir == tmp_path / "logic"

    def test_subdir_pack_with_own_logic(self, tmp_path):
        """A pack with its own logic/ uses it."""
        pack = tmp_path / "my_pack"
        (pack / "logic").mkdir(parents=True)
        pack_dir, logic_dir = resolve_pack_dirs(tmp_path, 'my_pack')
        assert pack_dir == pack
        assert logic_dir == pack / "logic"

    def test_subdir_pack_falls_back_to_root_logic(self, tmp_path):
        """A pack without logic/ falls back to the root logic/."""
        pack = tmp_path / "my_pack"
        pack.mkdir()
        (tmp_path / "logic").mkdir()
        pack_dir, logic_dir = resolve_pack_dirs(tmp_path, 'my_pack')
        assert pack_dir == pack
        assert logic_dir == tmp_path / "logic"
