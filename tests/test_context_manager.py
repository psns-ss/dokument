"""
Tests for app/logic/context_manager.py.

Tests thread-local DataContext class management for multi-pack support.
"""

import pytest
from pathlib import Path
from docx import Document

from logic.context_manager import (
    load_data_context_for_pack,
    get_data_context_class,
    invalidate_data_context,
)


class TestLoadDataContextForPack:
    """Tests for load_data_context_for_pack() - creates DataContext class per pack."""

    def test_creates_data_context_class(self, tmp_path):
        """load_data_context_for_pack should create a DataContext class."""
        # Create a descr.docx with a variable
        doc = Document()
        table = doc.add_table(rows=1, cols=4)
        table.rows[0].cells[0].text = "test_var"
        table.rows[0].cells[1].text = "Test variable"
        table.rows[0].cells[2].text = "✓"
        table.rows[0].cells[3].text = "text"

        descr_path = tmp_path / "descr.docx"
        doc.save(descr_path)

        lang_path = tmp_path / "language_rules.yaml"
        lang_path.write_text("")

        # Load the DataContext class
        DataContext = load_data_context_for_pack(descr_path, lang_path)

        # Should have the test_var field
        assert hasattr(DataContext, "test_var")
        instance = DataContext()
        assert instance.test_var == ""

    def test_reuses_class_for_same_pack(self, tmp_path):
        """load_data_context_for_pack should reuse class for same pack."""
        # Create a descr.docx
        doc = Document()
        table = doc.add_table(rows=1, cols=4)
        table.rows[0].cells[0].text = "var1"
        table.rows[0].cells[1].text = "Variable 1"
        table.rows[0].cells[2].text = "✓"
        table.rows[0].cells[3].text = "text"

        descr_path = tmp_path / "descr.docx"
        doc.save(descr_path)

        lang_path = tmp_path / "language_rules.yaml"
        lang_path.write_text("")

        # Load twice
        DataContext1 = load_data_context_for_pack(descr_path, lang_path)
        DataContext2 = load_data_context_for_pack(descr_path, lang_path)

        # Should be the same class
        assert DataContext1 is DataContext2

    def test_creates_different_class_for_different_pack(self, tmp_path):
        """load_data_context_for_pack should create different classes for different packs."""
        # Create pack A
        doc_a = Document()
        table_a = doc_a.add_table(rows=1, cols=4)
        table_a.rows[0].cells[0].text = "var_a"
        table_a.rows[0].cells[1].text = "Var A"
        table_a.rows[0].cells[2].text = "✓"
        table_a.rows[0].cells[3].text = "text"

        descr_path_a = tmp_path / "descr_a.docx"
        doc_a.save(descr_path_a)

        # Create pack B
        doc_b = Document()
        table_b = doc_b.add_table(rows=1, cols=4)
        table_b.rows[0].cells[0].text = "var_b"
        table_b.rows[0].cells[1].text = "Var B"
        table_b.rows[0].cells[2].text = "✓"
        table_b.rows[0].cells[3].text = "text"

        descr_path_b = tmp_path / "descr_b.docx"
        doc_b.save(descr_path_b)

        lang_path = tmp_path / "language_rules.yaml"
        lang_path.write_text("")

        # Load both
        DataContext_a = load_data_context_for_pack(descr_path_a, lang_path)
        DataContext_b = load_data_context_for_pack(descr_path_b, lang_path)

        # Should be different classes with different fields
        assert DataContext_a is not DataContext_b
        assert hasattr(DataContext_a, "var_a")
        assert hasattr(DataContext_b, "var_b")


class TestGetDataContextClass:
    """Tests for get_data_context_class() - retrieves current DataContext class."""

    def test_raises_error_when_not_initialized(self):
        """get_data_context_class should raise RuntimeError when not initialized."""
        # Invalidate any existing context
        invalidate_data_context()

        with pytest.raises(RuntimeError, match="DataContext not initialized"):
            get_data_context_class()

    def test_returns_loaded_class(self, tmp_path):
        """get_data_context_class should return the loaded class."""
        # Create a descr.docx
        doc = Document()
        table = doc.add_table(rows=1, cols=4)
        table.rows[0].cells[0].text = "test_var"
        table.rows[0].cells[1].text = "Test"
        table.rows[0].cells[2].text = "✓"
        table.rows[0].cells[3].text = "text"

        descr_path = tmp_path / "descr.docx"
        doc.save(descr_path)

        lang_path = tmp_path / "language_rules.yaml"
        lang_path.write_text("")

        # Load first
        DataContext = load_data_context_for_pack(descr_path, lang_path)

        # Then retrieve
        retrieved = get_data_context_class()

        assert retrieved is DataContext


class TestInvalidateDataContext:
    """Tests for invalidate_data_context() - clears thread-local cache."""

    def test_clears_cache(self, tmp_path):
        """invalidate_data_context should clear the cached class."""
        # Create a descr.docx
        doc = Document()
        table = doc.add_table(rows=1, cols=4)
        table.rows[0].cells[0].text = "test_var"
        table.rows[0].cells[1].text = "Test"
        table.rows[0].cells[2].text = "✓"
        table.rows[0].cells[3].text = "text"

        descr_path = tmp_path / "descr.docx"
        doc.save(descr_path)

        lang_path = tmp_path / "language_rules.yaml"
        lang_path.write_text("")

        # Load
        load_data_context_for_pack(descr_path, lang_path)

        # Invalidate
        invalidate_data_context()

        # Should raise error now
        with pytest.raises(RuntimeError, match="DataContext not initialized"):
            get_data_context_class()

    def test_allows_reload_after_invalidation(self, tmp_path):
        """Should be able to load a new context after invalidation."""
        # Create first descr.docx
        doc1 = Document()
        table1 = doc1.add_table(rows=1, cols=4)
        table1.rows[0].cells[0].text = "var1"
        table1.rows[0].cells[1].text = "Var 1"
        table1.rows[0].cells[2].text = "✓"
        table1.rows[0].cells[3].text = "text"

        descr_path1 = tmp_path / "descr1.docx"
        doc1.save(descr_path1)

        # Create second descr.docx
        doc2 = Document()
        table2 = doc2.add_table(rows=1, cols=4)
        table2.rows[0].cells[0].text = "var2"
        table2.rows[0].cells[1].text = "Var 2"
        table2.rows[0].cells[2].text = "✓"
        table2.rows[0].cells[3].text = "text"

        descr_path2 = tmp_path / "descr2.docx"
        doc2.save(descr_path2)

        lang_path = tmp_path / "language_rules.yaml"
        lang_path.write_text("")

        # Load first
        load_data_context_for_pack(descr_path1, lang_path)

        # Invalidate
        invalidate_data_context()

        # Load second
        DataContext2 = load_data_context_for_pack(descr_path2, lang_path)

        # Should have var2 field
        assert hasattr(DataContext2, "var2")
