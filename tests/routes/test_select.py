"""
Tests for app/routes/select.py - pack selection route.
"""

import pytest
from pathlib import Path
from docx import Document


def create_descr_with_description(path: Path, description: str):
    """Helper to create a descr.docx with a description paragraph."""
    doc = Document()
    if description:
        doc.add_paragraph(description)

    # Add a table
    table = doc.add_table(rows=1, cols=4)
    table.rows[0].cells[0].text = "var1"
    table.rows[0].cells[1].text = "Variable 1"
    table.rows[0].cells[2].text = "✓"
    table.rows[0].cells[3].text = "text"

    doc.save(path)


class TestSelectRoute:
    """Tests for the /select route."""

    def test_select_get_shows_form_multiple_packs(self, client, tmp_path):
        """GET /select should show the selection form when multiple packs exist."""
        # Create two packs to prevent auto-select
        create_descr_with_description(tmp_path / "descr.docx", "Root pack")

        subdir = tmp_path / "my_pack"
        subdir.mkdir()
        create_descr_with_description(subdir / "descr.docx", "My pack")

        # Temporarily modify cfg.templates_dir
        from config import cfg
        original_dir = cfg.templates_dir
        cfg.templates_dir = tmp_path

        try:
            response = client.get('/select')
            assert response.status_code == 200
            assert 'Выберите набор шаблонов' in response.data.decode('utf-8')
        finally:
            cfg.templates_dir = original_dir

    def test_select_get_single_pack_auto_selects(self, client, tmp_path):
        """GET /select with single pack should auto-select and redirect."""
        # Create root pack only
        create_descr_with_description(tmp_path / "descr.docx", "Root pack")

        from config import cfg
        original_dir = cfg.templates_dir
        cfg.templates_dir = tmp_path

        try:
            response = client.get('/select', follow_redirects=False)
            # Should redirect to /init
            assert response.status_code == 302
            assert '/init' in response.location
        finally:
            cfg.templates_dir = original_dir

    def test_select_post_switches_pack(self, client, tmp_path):
        """POST /select with selected pack should switch and redirect."""
        # Create two packs
        create_descr_with_description(tmp_path / "descr.docx", "Root pack")

        subdir = tmp_path / "my_pack"
        subdir.mkdir()
        create_descr_with_description(subdir / "descr.docx", "My pack")

        from config import cfg
        original_dir = cfg.templates_dir
        cfg.templates_dir = tmp_path

        try:
            response = client.post('/select', data={'selected_pack': 'my_pack'}, follow_redirects=False)
            assert response.status_code == 302
            assert '/init' in response.location

            # Check session was updated
            with client.session_transaction() as sess:
                assert sess.get('active_pack') == 'my_pack'
        finally:
            cfg.templates_dir = original_dir

    def test_select_post_invalid_pack(self, client, tmp_path):
        """POST /select with invalid pack should show error."""
        # Create root pack only
        create_descr_with_description(tmp_path / "descr.docx", "Root pack")

        from config import cfg
        original_dir = cfg.templates_dir
        cfg.templates_dir = tmp_path

        try:
            response = client.post('/select', data={'selected_pack': 'nonexistent'}, follow_redirects=False)
            # Should show form again (render_template returns 200)
            assert response.status_code == 200
        finally:
            cfg.templates_dir = original_dir

    def test_select_get_no_packs_shows_error(self, client, tmp_path):
        """GET /select with no packs should show error message."""
        from config import cfg
        original_dir = cfg.templates_dir
        cfg.templates_dir = tmp_path

        try:
            response = client.get('/select')
            assert response.status_code == 200
            assert b'No template packs found' in response.data
        finally:
            cfg.templates_dir = original_dir

    def test_select_shows_template_count(self, client, tmp_path):
        """Select page should show template count for each pack."""
        # Create two packs to prevent auto-select
        create_descr_with_description(tmp_path / "descr.docx", "Root pack")
        (tmp_path / "template1.docx").write_bytes(b'dummy')
        (tmp_path / "template2.docx").write_bytes(b'dummy')

        subdir = tmp_path / "my_pack"
        subdir.mkdir()
        create_descr_with_description(subdir / "descr.docx", "My pack")

        from config import cfg
        original_dir = cfg.templates_dir
        cfg.templates_dir = tmp_path

        try:
            response = client.get('/select')
            assert response.status_code == 200
            # Should show "2 шаблонов" (2 templates in Russian)
            assert '2' in response.data.decode('utf-8')
        finally:
            cfg.templates_dir = original_dir

    def test_select_shows_description(self, client, tmp_path):
        """Select page should show pack descriptions."""
        # Create two packs to prevent auto-select
        create_descr_with_description(tmp_path / "descr.docx", "This is my root pack description")

        subdir = tmp_path / "my_pack"
        subdir.mkdir()
        create_descr_with_description(subdir / "descr.docx", "My pack description")

        from config import cfg
        original_dir = cfg.templates_dir
        cfg.templates_dir = tmp_path

        try:
            response = client.get('/select')
            assert response.status_code == 200
            assert 'This is my root pack description' in response.data.decode('utf-8')
        finally:
            cfg.templates_dir = original_dir


class TestPackSwitching:
    """Tests for pack switching behavior."""

    def test_session_data_cleared_on_switch(self, client, tmp_path):
        """Session data should be cleared when switching packs."""
        # Create two packs
        create_descr_with_description(tmp_path / "descr.docx", "Root pack")

        subdir = tmp_path / "my_pack"
        subdir.mkdir()
        create_descr_with_description(subdir / "descr.docx", "My pack")

        from config import cfg
        original_dir = cfg.templates_dir
        cfg.templates_dir = tmp_path

        try:
            # First, select root pack
            client.post('/select', data={'selected_pack': 'root'})

            # Add some session data
            with client.session_transaction() as sess:
                sess['user_name'] = 'John'
                sess['user_age'] = '30'

            # Switch to my_pack
            response = client.post('/select', data={'selected_pack': 'my_pack'}, follow_redirects=False)
            assert response.status_code == 302

            # Check session was cleared (except system keys)
            with client.session_transaction() as sess:
                assert 'user_name' not in sess
                assert 'user_age' not in sess
                assert sess.get('active_pack') == 'my_pack'
        finally:
            cfg.templates_dir = original_dir
