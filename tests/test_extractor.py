"""
Tests for app/logic/extractor.py::variables_from_docx.

variables_from_docx extracts variable names from a DOCX file, normalizing every
docxtpl placeholder form to its base variable:
  {{var}} / {{var.case}} / {{p var}} / {{l var}} / {{t var}} / {{! var}} /
  {{p var[0]}} / {{# var}} / {{/ var}} (loop tags)
  {{var[key]}} (identifier subscript — a context lookup),
  {% if var ... %} / {%p if ... %} (conditions),
  {% for x in var %} (iterable; loop targets are suppressed as local)
It also scans every headerN/footerN part (not just the first nine).
"""

import pytest

from logic.extractor import variables_from_docx


class TestVariablesFromDocx:
    """Tests for the variables_from_docx function."""

    def test_plain_variables(self, fixtures_dir):
        """Plain {{var}} placeholders are extracted."""
        result = variables_from_docx(fixtures_dir / "template_simple.docx")
        assert {"user_name", "user_sex", "birth_year", "company_name"} <= result

    def test_dotted_inflection_normalizes_to_base(self, fixtures_dir):
        """{{var.case}} yields the base variable, not the dotted form."""
        result = variables_from_docx(fixtures_dir / "template_inflection.docx")
        assert "org_name" in result
        # the dotted form itself must NOT be reported as a variable
        assert "org_name.no" not in result
        assert "org_name.pre" not in result

    def test_tagged_subdoc_normalizes_to_base(self, fixtures_dir):
        """{{p var}} (rich/subdoc tag) yields the base variable."""
        result = variables_from_docx(fixtures_dir / "template_inflection.docx")
        assert "table_3" in result
        assert "university_historical" in result
        # the tagged form itself is not a variable name
        assert "p table_3" not in result

    def test_safe_and_subscript_forms(self, fixtures_dir):
        """{{! var}} (safe) and {{p var[0]}} (subscript) normalize to the base."""
        result = variables_from_docx(fixtures_dir / "template_inflection.docx")
        assert "user_name" in result            # from {{! user_name}}
        assert "university_historical" in result  # from {{p university_historical[0]}}

    def test_single_letter_variable(self, fixtures_dir):
        """A single-letter variable ({{f}}) is a variable, not a tag."""
        result = variables_from_docx(fixtures_dir / "template_inflection.docx")
        assert "f" in result

    def test_high_numbered_headers_are_scanned(self, fixtures_dir):
        """Variables in header parts beyond header9.xml are found (>=10 headers)."""
        result = variables_from_docx(fixtures_dir / "template_many_headers.docx")
        # sentinel lives only in the last (highest-numbered) header part
        assert "deep_header_var" in result
        assert "hdr_var" in result
        assert "body_var" in result

    def test_loop_tags_normalize_to_base(self, tmp_path):
        """{{# var}} / {{/ var}} (loop tags) yield the base variable."""
        from docx import Document
        doc = Document()
        doc.add_paragraph("{{# items}} row {{/ items}}")
        path = tmp_path / "loop_tags.docx"
        doc.save(str(path))

        result = variables_from_docx(path)
        assert "items" in result

    def test_if_block_variable_is_found(self, tmp_path):
        """{%p if var ... %} — the condition's variable is extracted;
        the string literal is not."""
        from docx import Document
        doc = Document()
        doc.add_paragraph("{%p if signature == 'true' %}")
        doc.add_paragraph("С уважением,")
        doc.add_paragraph("{%p endif %}")
        path = tmp_path / "if_block.docx"
        doc.save(str(path))

        result = variables_from_docx(path)
        assert "signature" in result
        assert "true" not in result

    def test_plain_if_block_with_curly_quotes(self, tmp_path):
        """{% if sex == ‘Мужской’ %} — stock {% form and curly quotes
        (docxtpl's clean_tags normalizes them before Jinja)."""
        from docx import Document
        doc = Document()
        doc.add_paragraph("Уважаем{% if sex == ‘Мужской’ %}ый{% else %}ая{% endif %} {{user_name}}!")
        path = tmp_path / "curly_quotes.docx"
        doc.save(str(path))

        result = variables_from_docx(path)
        assert "sex" in result
        assert "user_name" in result
        assert "Мужской" not in result

    def test_for_block_iterable_found_target_suppressed(self, tmp_path):
        """{%p for u in university %} — the iterable is a variable, the
        loop target is local: {{ u.full }} and u.year|int yield no "u"."""
        from docx import Document
        doc = Document()
        doc.add_paragraph("{%p for u in university %}")
        doc.add_paragraph("University: {{ u.full }}")
        doc.add_paragraph("{%p if u.year|int > 2019 %}")
        doc.add_paragraph("Graduated recently!")
        doc.add_paragraph("{%p endif %}")
        doc.add_paragraph("{%p endfor %}")
        path = tmp_path / "for_block.docx"
        doc.save(str(path))

        result = variables_from_docx(path)
        assert "university" in result
        assert "u" not in result
        assert "int" not in result  # filter name, not a variable

    def test_for_tuple_targets_are_local(self, tmp_path):
        """{% for a, b in pairs %} — all tuple targets are loop-local."""
        from docx import Document
        doc = Document()
        doc.add_paragraph("{% for a, b in pairs %}{{ a }}:{{ b }}{% endfor %}")
        path = tmp_path / "for_tuple.docx"
        doc.save(str(path))

        result = variables_from_docx(path)
        assert "pairs" in result
        assert "a" not in result
        assert "b" not in result

    def test_subscript_identifier_is_a_reference(self, tmp_path):
        """{{p var[key]}} — an identifier subscript is a context lookup
        (Jinja resolves `key` in the data), so both names are reported."""
        from docx import Document
        doc = Document()
        doc.add_paragraph("{{p university_historical[moscow_state_university] }}")
        path = tmp_path / "subscript_ident.docx"
        doc.save(str(path))

        result = variables_from_docx(path)
        assert "university_historical" in result
        assert "moscow_state_university" in result

    def test_subscript_index_and_string_key_not_reported(self, tmp_path):
        """{{p var[0]}} / {{var['key']}} — numeric indices and string
        keys are not variables."""
        from docx import Document
        doc = Document()
        doc.add_paragraph("{{p university_historical[1] }}")
        doc.add_paragraph("{{ org['name'] }}")
        path = tmp_path / "subscript_other.docx"
        doc.save(str(path))

        result = variables_from_docx(path)
        assert "university_historical" in result
        assert "org" in result
        assert "name" not in result

    def test_variable_after_textbox_in_same_paragraph(self, tmp_path):
        """X1: a textbox nests <w:p> inside <w:p> (via <w:txbxContent>), so a
        var after the textbox in the same outer paragraph must still be found
        (the old non-greedy split stopped at the textbox's first </w:p>)."""
        from docx import Document
        from docx.oxml import parse_xml
        doc = Document()
        p_xml = (
            '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
            ' xmlns:v="urn:schemas-microsoft-com:vml">'
            "<w:r><w:t>до {{before_var}}</w:t></w:r>"
            "<w:r><w:pict>"
            '<v:shape id="tb1" type="#_x0000_t20">'
            "<v:textbox><w:txbxContent>"
            "<w:p><w:r><w:t>{{tb_var}}</w:t></w:r></w:p>"
            "</w:txbxContent></v:textbox>"
            "</v:shape>"
            "</w:pict></w:r>"
            "<w:r><w:t>после {{after_var}}</w:t></w:r>"
            "</w:p>"
        )
        doc.element.body.append(parse_xml(p_xml))
        path = tmp_path / "textbox.docx"
        doc.save(str(path))

        result = variables_from_docx(path)
        assert "before_var" in result
        assert "tb_var" in result
        assert "after_var" in result  # the regression: dropped by the old split

    def test_variable_after_self_closing_paragraph(self, tmp_path):
        """A self-closing <w:p .../> (empty paragraph Word emits) must not
        count as an open tag: the depth walk would stick above 0 and drop
        every var after the first empty paragraph in the document."""
        from docx import Document
        from docx.oxml import parse_xml
        W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        doc = Document()
        doc.element.body.append(parse_xml(
            f'<w:p {W}><w:r><w:t>{{{{before_var}}}}</w:t></w:r></w:p>'))
        doc.element.body.append(parse_xml(f'<w:p {W}/>'))  # empty paragraph
        doc.element.body.append(parse_xml(
            f'<w:p {W}><w:r><w:t>{{{{after_var}}}}</w:t></w:r></w:p>'))
        path = tmp_path / "selfclosing.docx"
        doc.save(str(path))

        result = variables_from_docx(path)
        assert "before_var" in result
        assert "after_var" in result  # the regression: depth stuck, var dropped

    def test_missing_file_raises(self, fixtures_dir):
        """A non-existent path raises (zipfile behaviour, relied on by debug.py)."""
        with pytest.raises(FileNotFoundError):
            variables_from_docx(fixtures_dir / "does_not_exist.docx")
