"""
Tests for logic/university.py - subdoc generation functionality.

This tests the university_historical() function which generates subdocuments
for selected universities based on bookmarks in university.docx.
"""

import pytest
from pathlib import Path
import sys
from unittest.mock import MagicMock

# Add project root to path for imports
project_root = Path(__file__).parent.parent.parent
sys.path.insert(0, str(project_root))

from docxtpl import DocxTemplate
from docx import Document

# Clear any cached logic.university module from previous test imports
sys.modules.pop('logic.university', None)

# Mock the logic.matcher module that university.py imports
matcher_mock = MagicMock()
matcher_mock.find_matches = lambda raw_text, yaml_dict, keys=("short", "full"), use_fuzzy=True, threshold=80: []
sys.modules['logic.matcher'] = matcher_mock

# Mock logic.subdoc_helper to use the real app.logic.subdoc_helper
import logic.subdoc_helper as real_subdoc_helper
subdoc_helper_mock = MagicMock()
subdoc_helper_mock.create_subdoc_from_bookmark = real_subdoc_helper.create_subdoc_from_bookmark
sys.modules['logic.subdoc_helper'] = subdoc_helper_mock

# Now we can import the actual university module with real subdoc support
from logic.university import university_historical


class TestUniversityHistoricalSubdoc:
    """Tests for the university_historical subdoc generation function."""

    def test_university_historical_single_university(self, project_root):
        """Test generating subdoc for a single university."""
        template_path = project_root / "templates" / "template_3.docx"
        tpl = DocxTemplate(str(template_path))

        kwargs = {"university": [{"key": "moscow_state_university"}]}
        result = university_historical(tpl, **kwargs)

        assert isinstance(result, dict)
        assert "moscow_state_university" in result

    def test_university_historical_multiple_universities(self, project_root):
        """Test generating subdocs for multiple universities."""
        template_path = project_root / "templates" / "template_3.docx"
        tpl = DocxTemplate(str(template_path))

        kwargs = {
            "university": [
                {"key": "moscow_state_university"},
                {"key": "bauman_mgtu"},
            ]
        }

        result = university_historical(tpl, **kwargs)

        assert isinstance(result, dict)
        assert len(result) == 2
        assert "moscow_state_university" in result
        assert "bauman_mgtu" in result

    def test_university_historical_empty_list(self, project_root):
        """Test that empty university list returns empty dict."""
        template_path = project_root / "templates" / "template_3.docx"
        tpl = DocxTemplate(str(template_path))

        kwargs = {"university": []}
        result = university_historical(tpl, **kwargs)

        assert result == {}

    def test_university_historical_subdoc_has_content(self, project_root):
        """Test that generated subdocs contain actual content."""
        template_path = project_root / "templates" / "template_3.docx"
        tpl = DocxTemplate(str(template_path))

        kwargs = {"university": [{"key": "moscow_state_university"}]}
        subdoc_dict = university_historical(tpl, **kwargs)

        assert "moscow_state_university" in subdoc_dict
        subdoc = subdoc_dict["moscow_state_university"]
        assert subdoc is not None

        # Extract content from subdoc
        if hasattr(subdoc, '_subdoc'):
            subdoc_buffer = subdoc._subdoc
            if hasattr(subdoc_buffer, 'getvalue'):
                subdoc_buffer.seek(0)
                subdoc_doc = Document(subdoc_buffer)
                all_text = "".join([p.text for p in subdoc_doc.paragraphs])
                # МГУ was founded in 1755
                assert "1755" in all_text or len(all_text) > 0

    def test_university_historical_full_render(self, project_root, tmp_path):
        """Test full rendering flow with university_historical and university list."""
        template_path = project_root / "templates" / "template_3.docx"
        tpl = DocxTemplate(str(template_path))

        university_data = [
            {"key": "moscow_state_university", "full": "Московский государственный университет", "short": "МГУ", "year": "2020"},
            {"key": "bauman_mgtu", "full": "МГТУ им. Баумана", "short": "МГТУ", "year": "2021"},
        ]

        subdoc_dict = university_historical(tpl, university=university_data)

        context = {
            "university": university_data,
            "university_historical": subdoc_dict,
            "user_name": "Иван Иванов",
            "sex": "Мужской",
        }

        output_path = tmp_path / "full_render_output.docx"
        tpl.render(context)
        tpl.save(str(output_path))

        assert output_path.exists()
        assert output_path.stat().st_size > 0

        output_doc = Document(str(output_path))
        all_text = "".join([p.text for p in output_doc.paragraphs])

        assert "Иван Иванов" in all_text
        assert "1755" in all_text
