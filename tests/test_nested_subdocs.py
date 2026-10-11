"""
Tests for nested subdoc rendering: bookmark-within-bookmark + table preservation.

Verifies cascading lazy evaluation via DeferredSubdoc:
  nested_base.docx  --{{p level1}}-->  nested_library.docx (bookmark "level1_content")
                                                   contains table column + {{p table_3}}
  nested_source.docx --table 0-->  placeholders {{user_name}}, {{user_age}}, {{sex}}

Flow:
  1. Render base template with DeferredSubdoc for level1 and table_3
  2. level1 extracts bookmark from library.docx (table column + {{p table_3}})
  3. Bookmark content rendered with table_3 as DeferredSubdoc
  4. table_3 extracts table 0 from source.docx
  5. Table placeholders sandwich-rendered with actual values
  6. Result propagates: table → bookmark subdoc → base doc
"""

import pytest
from pathlib import Path
import sys
import zipfile

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

from docx import Document
from docxtpl import DocxTemplate

from logic.renderer import render_templates
from logic.subdoc_helper import (
    DeferredSubdoc,
    create_subdoc_from_bookmark,
    extract_table_as_subdoc,
)
from logic.context import make_context_class


def extract_docx_text(doc_path: Path) -> str:
    """Extract all text from a DOCX (paragraphs + table cells)."""
    doc = Document(str(doc_path))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.append("|".join(c.text for c in row.cells))
    return "\n".join(parts)


def get_raw_xml_text(doc_path: Path) -> str:
    """Read raw document.xml and decode as UTF-8."""
    with zipfile.ZipFile(str(doc_path), 'r') as z:
        return z.read('word/document.xml').decode('utf-8')


class TestNestedSubdocs:
    """Nested subdoc rendering with bookmark extraction and table preservation."""

    def test_nested_subdoc_with_table(self, fixtures_dir, tmp_path):
        """
        base.docx has {{p level1}} which extracts a bookmark from library.docx.
        The bookmark contains a table column with {{user_name}} and {{p table_3}}.
        table_3 extracts table 0 from source.docx with {{user_name}}, {{user_age}}, {{sex}}.

        Verify:
          - user_name rendered at least twice (column text + table cell)
          - All placeholders resolved (no unrendered Jinja tags)
          - Tables preserved in output
        """
        base_tpl = fixtures_dir / "nested_base.docx"
        library_doc = fixtures_dir / "nested_library.docx"
        source_doc = fixtures_dir / "nested_source.docx"

        user_name = "TestUser Nested"
        context_data = {
            'user_name': user_name,
            'user_age': '42',
            'sex': 'Male',
            'table_1': '',
            'table_2': '',
        }

        # --- Subdoc functions (mirror user plugins) ---

        def level1_subdoc(tpl, **ctx):
            """Outer: extract bookmark from library.docx."""
            return create_subdoc_from_bookmark(
                tpl,
                bookmark_name="level1_content",
                source_path=library_doc,
                render_context=ctx,
            )

        def table_3_subdoc(tpl, **ctx):
            """Inner: extract table 0 from source.docx."""
            source = Document(str(source_doc))
            return extract_table_as_subdoc(
                tpl, source, 0,
                context_data=ctx,
                render_placeholders=True,
            )

        # --- Render via real framework ---
        tpl = DocxTemplate(str(base_tpl))
        data_dict = dict(context_data)
        data_dict['level1'] = DeferredSubdoc(level1_subdoc, tpl, data_dict, name='level1')
        data_dict['table_3'] = DeferredSubdoc(table_3_subdoc, tpl, data_dict, name='table_3')

        tpl.render(data_dict)

        out_path = tmp_path / "nested_result.docx"
        tpl.save(str(out_path))

        # --- Verify ---
        raw = get_raw_xml_text(out_path)

        # user_name should appear at least twice: column text + table cell
        count = raw.count(user_name)
        assert count >= 2, f"user_name '{user_name}' appears {count}x, expected >= 2"

        # No unrendered placeholders
        for tag in ['{{user_name}}', '{{user_age}}', '{{sex}}', '{{p table_3}}']:
            assert tag not in raw, f"Unrendered placeholder: {tag}"

        # Tables preserved
        result_doc = Document(str(out_path))
        assert len(result_doc.tables) >= 1, "Expected at least 1 table in output"

    def test_nested_subdoc_via_render_templates(self, fixtures_dir, tmp_path):
        """
        Same nested flow but through render_templates() (the real app entry point).
        Uses make_context_class for DataContext, subdoc_funcs dict.
        """
        # Build DataContext from project's descr.docx for field structure
        templates_dir = project_root / "templates"
        ContextClass = make_context_class(
            templates_dir / "descr.docx",
            templates_dir / "language_rules.yaml",
        )

        context = ContextClass(
            user_name="FrameworkTestUser",
            user_age="30",
            sex="Female",
        )

        store_dir = tmp_path / "session"
        store_dir.mkdir()

        # Subdoc functions passed to render_templates
        def level1_fn(tpl, **ctx):
            return create_subdoc_from_bookmark(
                tpl, "level1_content", fixtures_dir / "nested_library.docx",
                render_context=ctx,
            )

        def table_3_fn(tpl, **ctx):
            source = Document(str(fixtures_dir / "nested_source.docx"))
            return extract_table_as_subdoc(tpl, source, 0, ctx, True)

        output_files = render_templates(
            templates=[fixtures_dir / "nested_base.docx"],
            context=context,
            store_dir=store_dir,
            subdoc_funcs={
                'level1': level1_fn,
                'table_3': table_3_fn,
            },
        )

        assert len(output_files) == 1
        raw = get_raw_xml_text(output_files[0])

        count = raw.count("FrameworkTestUser")
        assert count >= 2, f"user_name appears {count}x, expected >= 2"

        for tag in ['{{user_name}}', '{{user_age}}', '{{sex}}', '{{p table_3}}']:
            assert tag not in raw, f"Unrendered placeholder: {tag}"

        result_doc = Document(str(output_files[0]))
        assert len(result_doc.tables) >= 1


class TestNestedSubdocEdgeCases:
    """Edge cases for nested subdoc rendering."""

    def test_inner_subdoc_table_out_of_range(self, fixtures_dir, tmp_path):
        """When inner subdoc references a non-existent table index, it returns empty."""
        tpl = DocxTemplate(str(fixtures_dir / "nested_base.docx"))

        def level1_fn(tpl, **ctx):
            return create_subdoc_from_bookmark(
                tpl, "level1_content", fixtures_dir / "nested_library.docx",
                render_context=ctx,
            )

        def bad_table_fn(tpl, **ctx):
            source = Document(str(fixtures_dir / "nested_source.docx"))
            # Request table index 999 — out of range
            return extract_table_as_subdoc(tpl, source, 999, ctx, True)

        data_dict = {
            'user_name': 'EdgeCaseUser',
            'user_age': '99',
            'sex': 'Other',
            'table_1': '',
            'table_2': '',
        }
        data_dict['level1'] = DeferredSubdoc(level1_fn, tpl, data_dict, name='level1')
        data_dict['table_3'] = DeferredSubdoc(bad_table_fn, tpl, data_dict, name='table_3')

        tpl.render(data_dict)

        out_path = tmp_path / "edge_result.docx"
        tpl.save(str(out_path))

        # Should still render — inner subdoc returns empty, outer continues
        raw = get_raw_xml_text(out_path)
        assert "EdgeCaseUser" in raw

    def test_bookmark_not_found_graceful(self, fixtures_dir, tmp_path):
        """When bookmark name doesn't exist in library, subdoc returns empty."""
        tpl = DocxTemplate(str(fixtures_dir / "nested_base.docx"))

        def missing_bm_fn(tpl, **ctx):
            return create_subdoc_from_bookmark(
                tpl, "nonexistent_bookmark", fixtures_dir / "nested_library.docx",
                render_context=ctx,
            )

        data_dict = {
            'user_name': 'MissingBMUser',
            'user_age': '1',
            'sex': 'X',
            'table_1': '',
            'table_2': '',
        }
        data_dict['level1'] = DeferredSubdoc(missing_bm_fn, tpl, data_dict, name='level1')
        data_dict['table_3'] = DeferredSubdoc(
            lambda tpl, **ctx: tpl.new_subdoc(), tpl, data_dict, name='table_3'
        )

        tpl.render(data_dict)

        out_path = tmp_path / "missing_bm_result.docx"
        tpl.save(str(out_path))

        # Should still render — missing bookmark → empty subdoc
        raw = get_raw_xml_text(out_path)
        assert "MissingBMUser" in raw
