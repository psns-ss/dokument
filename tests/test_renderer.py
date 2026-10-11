"""
Tests for app/logic/renderer.py

Tests the core document rendering pipeline including variable substitution
and subdocument rendering. Uses actual templates from templates/ directory.

Content verification uses unique strings to avoid false positives from
Word's XML structure quirks (text split across elements, etc.).
"""

import pytest
from pathlib import Path
import sys
from unittest.mock import MagicMock

# Add project root to path for imports
project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

from docx import Document
from docxtpl import DocxTemplate

from logic.renderer import render_templates
from logic.context import make_context_class


def extract_docx_text(doc_path: Path) -> str:
    """
    Extract all text from a DOCX file.

    Word's XML structure can split text across multiple elements,
    so we concatenate all paragraph text for reliable searching.
    """
    doc = Document(str(doc_path))
    paragraphs = [p.text for p in doc.paragraphs]
    return "\n".join(paragraphs)


class TestRenderTemplatesSimple:
    """Tests for basic template rendering with template_1.docx."""

    def test_render_template_1_basic_vars(self, project_root, tmp_path):
        """
        Test rendering template_1 with basic user variables.

        template_1 contains:
        - {{user_name}} - rendered as "QwertyXyz123"
        - {{user_age}} - computed field, rendered as "30"
        - {{sex}} - rendered as "Мужской"
        - {{birth_year}} in header

        We search for unique strings to avoid false positives.
        """
        templates_dir = project_root / "templates"

        # Create context class from descr.docx
        ContextClass = make_context_class(
            templates_dir / "descr.docx",
            templates_dir / "language_rules.yaml"
        )

        # Create context with unique test values
        context = ContextClass(
            user_name="QwertyXyz123",  # Unique string for verification
            user_age="30",
            sex="Мужской",
        )

        store_dir = tmp_path / "test_session"
        store_dir.mkdir(parents=True, exist_ok=True)

        # Render template_1
        output_files = render_templates(
            templates=[templates_dir / "template_1.docx"],
            context=context,
            store_dir=store_dir,
        )

        assert len(output_files) == 1
        assert output_files[0].exists()

        # Verify content - search for unique test string
        text = extract_docx_text(output_files[0])
        assert "QwertyXyz123" in text, f"Expected user_name not found in: {text[:500]}"
        assert "30" in text, f"Expected user_age not found in: {text[:500]}"



class TestRenderTemplatesAppeal:
    """Tests for template_2.docx with appeal computation."""

    def test_render_template_2_with_appeal_male(self, project_root, tmp_path):
        """
        Test rendering template_2 with computed appeal (male).

        template_2 contains:
        - {{appeal}} - computed from sex (should be "ый" for male)
        - {{user_name}} - "TestUserMale789"
        - {{birth_year}} - computed from age

        The phrase "Уважаемый TestUserMale789" should appear.
        """
        templates_dir = project_root / "templates"

        ContextClass = make_context_class(
            templates_dir / "descr.docx",
            templates_dir / "language_rules.yaml"
        )

        # Pre-compute appeal using the actual compute function
        from logic.appeal import compute_appeal
        appeal_value = compute_appeal(sex="Мужской")

        context = ContextClass(
            user_name="TestUserMale789",
            sex="Мужской",
            user_age="35",
            appeal=appeal_value,  # Pre-computed
        )

        store_dir = tmp_path / "test_session"
        store_dir.mkdir(parents=True, exist_ok=True)

        output_files = render_templates(
            templates=[templates_dir / "template_2.docx"],
            context=context,
            store_dir=store_dir,
        )

        assert len(output_files) == 1

        text = extract_docx_text(output_files[0])
        # "Уважаемый" is the male formal appeal (Уважаем + ый)
        assert "Уважаемый" in text, f"Expected male appeal in: {text[:500]}"
        assert "TestUserMale789" in text

    def test_render_template_2_with_appeal_female(self, project_root, tmp_path):
        """
        Test rendering template_2 with computed appeal (female).

        The phrase "Уважаемая TestUserFemale000" should appear.
        """
        templates_dir = project_root / "templates"

        ContextClass = make_context_class(
            templates_dir / "descr.docx",
            templates_dir / "language_rules.yaml"
        )

        # Pre-compute appeal using the actual compute function
        from logic.appeal import compute_appeal
        appeal_value = compute_appeal(sex="Женский")

        context = ContextClass(
            user_name="TestUserFemale000",
            sex="Женский",
            user_age="28",
            appeal=appeal_value,  # Pre-computed
        )

        store_dir = tmp_path / "test_session"
        store_dir.mkdir(parents=True, exist_ok=True)

        output_files = render_templates(
            templates=[templates_dir / "template_2.docx"],
            context=context,
            store_dir=store_dir,
        )

        assert len(output_files) == 1

        text = extract_docx_text(output_files[0])
        # "Уважаемая" is the female formal appeal (Уважаем + ая)
        assert "Уважаемая" in text, f"Expected female appeal in: {text[:500]}"
        assert "TestUserFemale000" in text


class TestRenderTemplatesUniversity:
    """Tests for template_3.docx with university loop and subdocs."""

    def test_render_template_3_with_university_list(self, project_root, tmp_path):
        """
        Test rendering template_3 with university list loop.

        template_3 contains:
        - Loop over university list with {{ u.full }} and {{ u.short }}
        - Subdoc placeholders for university_historical

        We verify the university names appear in output.
        """
        templates_dir = project_root / "templates"

        ContextClass = make_context_class(
            templates_dir / "descr.docx",
            templates_dir / "language_rules.yaml"
        )

        # Create university list as compute_university would
        university_list = [
            {
                "key": "moscow_state_university",
                "short": "МГУ",
                "full": "Московский государственный университет",
                "year": "2020",
            },
            {
                "key": "bauman_mgtu",
                "short": "МГТУ",
                "full": "МГТУ им. Баумана",
                "year": "2021",
            },
        ]

        context = ContextClass(
            user_name="UniversityStudent999",
            sex="Мужской",
            university=university_list,
        )

        store_dir = tmp_path / "test_session"
        store_dir.mkdir(parents=True, exist_ok=True)

        output_files = render_templates(
            templates=[templates_dir / "template_3.docx"],
            context=context,
            store_dir=store_dir,
        )

        assert len(output_files) == 1

        text = extract_docx_text(output_files[0])
        # Verify university names from the loop
        assert "Московский государственный университет" in text
        assert "МГТУ им. Баумана" in text
        assert "UniversityStudent999" in text

    def test_render_template_3_with_subdocs(self, project_root, tmp_path):
        """
        Test rendering template_3 with university subdocs.

        This tests the full subdoc pipeline:
        1. user selects universities
        2. university_historical() creates subdocs from bookmarks
        3. subdocs are rendered into template

        We verify МГУ's founding year (1755) appears in output.
        """
        templates_dir = project_root / "templates"

        ContextClass = make_context_class(
            templates_dir / "descr.docx",
            templates_dir / "language_rules.yaml"
        )

        # Import user logic for subdoc generation
        from logic.university import university_historical

        university_list = [
            {"key": "moscow_state_university", "short": "МГУ", "full": "МГУ Full", "year": "2020"},
        ]

        context = ContextClass(
            user_name="SubdocTestUser",
            sex="Женский",
            university=university_list,
        )

        store_dir = tmp_path / "test_session"
        store_dir.mkdir(parents=True, exist_ok=True)

        output_files = render_templates(
            templates=[templates_dir / "template_3.docx"],
            context=context,
            store_dir=store_dir,
            subdoc_funcs={"university_historical": university_historical},
        )

        assert len(output_files) == 1

        text = extract_docx_text(output_files[0])
        # МГУ was founded in 1755 - this should appear from the subdoc
        assert "1755" in text, f"Expected МГУ founding year 1755 in: {text[:1000]}"
        assert "SubdocTestUser" in text

    def test_render_template_3_full_workflow(self, project_root, tmp_path):
        """
        Test complete template_3 workflow with multiple universities and subdocs.

        This simulates the actual production flow:
        1. User fills form with name, sex
        2. User selects universities with years
        3. Both university list loop and subdocs are rendered
        """
        templates_dir = project_root / "templates"

        ContextClass = make_context_class(
            templates_dir / "descr.docx",
            templates_dir / "language_rules.yaml"
        )

        from logic.university import university_historical

        university_list = [
            {"key": "moscow_state_university", "short": "МГУ", "full": "Московский государственный университет", "year": "2019"},
            {"key": "bauman_mgtu", "short": "МГТУ", "full": "МГТУ им. Баумана", "year": "2021"},
        ]

        context = ContextClass(
            user_name="CompleteWorkflowTest",
            sex="Мужской",
            university=university_list,
        )

        store_dir = tmp_path / "test_session"
        store_dir.mkdir(parents=True, exist_ok=True)

        output_files = render_templates(
            templates=[templates_dir / "template_3.docx"],
            context=context,
            store_dir=store_dir,
            subdoc_funcs={"university_historical": university_historical},
        )

        assert len(output_files) == 1

        text = extract_docx_text(output_files[0])

        # Verify user data
        assert "CompleteWorkflowTest" in text

        # Verify university list loop rendered
        assert "Московский государственный университет" in text
        assert "МГТУ им. Баумана" in text

        # Verify subdoc content (1755 from МГУ historical)
        assert "1755" in text, f"Expected subdoc content in: {text[:1000]}"


class TestRenderTemplatesMultiple:
    """Tests for rendering multiple templates in one call."""

    def test_render_multiple_templates(self, project_root, tmp_path):
        """
        Test rendering multiple templates with same context.
        """
        templates_dir = project_root / "templates"

        ContextClass = make_context_class(
            templates_dir / "descr.docx",
            templates_dir / "language_rules.yaml"
        )

        context = ContextClass(
            user_name="MultiTemplateTest",
            sex="Мужской",
            user_age="40",
        )

        store_dir = tmp_path / "test_session"
        store_dir.mkdir(parents=True, exist_ok=True)

        output_files = render_templates(
            templates=[
                templates_dir / "template_1.docx",
                templates_dir / "template_2.docx",
            ],
            context=context,
            store_dir=store_dir,
        )

        assert len(output_files) == 2

        # Each output should contain the user name
        for output_file in output_files:
            text = extract_docx_text(output_file)
            assert "MultiTemplateTest" in text


class TestRenderTemplatesPartialFailure:
    """errors_out capture: a failing template is skipped + reported, good ones render."""

    def test_partial_failure_reports_errors(self, project_root, tmp_path):
        """A failing template is skipped and reported in errors_out; the good one still renders."""
        templates_dir = project_root / "templates"
        ContextClass = make_context_class(
            templates_dir / "descr.docx",
            templates_dir / "language_rules.yaml"
        )
        context = ContextClass(user_name="PartialFail789", user_age="40", sex="Мужской")

        store_dir = tmp_path / "test_session"
        store_dir.mkdir(parents=True, exist_ok=True)

        good = templates_dir / "template_1.docx"
        bad = templates_dir / "does_not_exist_999.docx"  # fails at DocxTemplate()

        errors_out = []
        output_files = render_templates(
            templates=[good, bad],
            context=context,
            store_dir=store_dir,
            errors_out=errors_out,
        )

        # Good template rendered; bad one skipped.
        assert len(output_files) == 1
        assert output_files[0].exists()
        # Failure captured with the bad template's name + a non-empty reason.
        assert len(errors_out) == 1
        assert errors_out[0][0] == bad.name
        assert errors_out[0][1]

    def test_no_errors_out_keeps_legacy_behavior(self, project_root, tmp_path):
        """Without errors_out, a failing template is skipped (empty outputs, no crash)."""
        templates_dir = project_root / "templates"
        ContextClass = make_context_class(
            templates_dir / "descr.docx",
            templates_dir / "language_rules.yaml"
        )
        context = ContextClass(user_name="LegacyCheck000", user_age="22", sex="Женский")

        store_dir = tmp_path / "test_session"
        store_dir.mkdir(parents=True, exist_ok=True)

        bad = templates_dir / "does_not_exist_999.docx"
        output_files = render_templates(
            templates=[bad],
            context=context,
            store_dir=store_dir,
        )

        assert output_files == []
