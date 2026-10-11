#!/usr/bin/env python3
"""
Create test fixture DOCX files.

Run this script to generate the sample DOCX files needed for testing.
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

# Get the directory where this script is located
script_dir = Path(__file__).parent


def create_sample_descr():
    """
    Create a sample descr.docx with test variable definitions.

    Table columns: Variable Name | Description | Show in Form | Type | Options/Path
    Note: No header row - parser expects data rows only
    """
    doc = Document()
    doc.add_heading('Test Variable Descriptions', 0)

    # Create table with variable definitions (no header row - parser expects data)
    # Data rows only
    test_vars = [
        ['user_name', 'Name of the user', '✓', 'text', ''],
        ['user_sex', 'Sex of the user', '✓', 'enum', 'sex_options'],
        ['birth_year', 'Year of birth', '✓', 'text', ''],
        ['hidden_var', 'Internal variable', '❌', 'text', ''],
        ['company_name', 'Name of company', '✓', 'match', 'company_1'],
    ]

    table = doc.add_table(rows=len(test_vars), cols=5)
    table.style = 'Table Grid'

    for i, var_data in enumerate(test_vars):
        for j, text in enumerate(var_data):
            table.rows[i].cells[j].text = text

    # Save the document
    output_path = script_dir / 'descr.docx'
    doc.save(output_path)
    print(f"Created: {output_path}")


def create_simple_template():
    """
    Create a simple template with basic variables.
    """
    doc = Document()
    doc.add_heading('Test Document', 0)

    doc.add_paragraph('Hello, {{user_name}}!')
    doc.add_paragraph('Your sex is: {{user_sex}}')
    doc.add_paragraph('Birth year: {{birth_year}}')
    doc.add_paragraph('Company: {{company_name}}')

    output_path = script_dir / 'template_simple.docx'
    doc.save(output_path)
    print(f"Created: {output_path}")


def create_template_with_header():
    """
    Create a template with variables in header.
    """
    doc = Document()

    # Add variable in header
    header = doc.sections[0].header
    header_paragraph = header.paragraphs[0] if doc.sections[0].header.paragraphs else header.add_paragraph()
    header_paragraph.text = 'Document for: {{user_name}}'

    # Add content
    doc.add_heading('Main Content', 1)
    doc.add_paragraph('This is the main body with {{birth_year}}.')

    output_path = script_dir / 'template_with_header.docx'
    doc.save(output_path)
    print(f"Created: {output_path}")


def create_template_with_table():
    """
    Create a template with a table.
    """
    doc = Document()
    doc.add_heading('Document with Table', 0)

    doc.add_paragraph('User: {{user_name}}')

    # Add a table
    table = doc.add_table(rows=3, cols=3)
    table.style = 'Table Grid'

    # Header row
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Column 1'
    hdr_cells[1].text = 'Column 2'
    hdr_cells[2].text = 'Column 3'

    # Data rows
    table.rows[1].cells[0].text = 'Data 1'
    table.rows[1].cells[1].text = 'Data 2'
    table.rows[1].cells[2].text = 'Data 3'

    table.rows[2].cells[0].text = 'Data 4'
    table.rows[2].cells[1].text = 'Data 5'
    table.rows[2].cells[2].text = 'Data 6'

    output_path = script_dir / 'template_with_table.docx'
    doc.save(output_path)
    print(f"Created: {output_path}")


def create_template_with_inflection():
    """
    Create a template exercising every docxtpl placeholder form the extractor
    must normalize to its base variable:
      plain {{user_name}}, dotted/inflection {{org_name.no}},
      tagged subdoc {{p table_3}}, safe {{! user_name}},
      subscript {{p university_historical[0]}}, single-letter {{f}}.
    """
    doc = Document()
    doc.add_heading('Forms Template', 0)
    doc.add_paragraph('Name: {{user_name}}')
    doc.add_paragraph('Org (nom): {{org_name.no}}')
    doc.add_paragraph('Org (prep): {{org_name.pre}}')
    doc.add_paragraph('Table: {{p table_3}}')
    doc.add_paragraph('Safe name: {{! user_name}}')
    doc.add_paragraph('History: {{p university_historical[0]}}')
    doc.add_paragraph('Single letter: {{f}}')

    output_path = script_dir / 'template_inflection.docx'
    doc.save(output_path)
    print(f"Created: {output_path}")


def create_template_many_headers():
    """
    Create a template with >=10 header parts (one per section) so a header
    beyond header9.xml is present. The old extractor only scanned
    header1..header9 and silently dropped the rest. A sentinel variable lives
    ONLY in the last section's header, so finding it proves the high-numbered
    header parts are scanned.
    """
    doc = Document()
    doc.add_paragraph('Body: {{body_var}}')

    # Section 0 header
    doc.sections[0].header.is_linked_to_previous = False
    doc.sections[0].header.paragraphs[0].text = 'Header: {{hdr_var}}'

    # 12 more sections -> header2.xml .. header13.xml (13 header parts total)
    for i in range(1, 13):
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        section.header.is_linked_to_previous = False
        section.header.paragraphs[0].text = 'Header: {{hdr_var}}'
        if i == 12:  # last section -> highest-numbered header part
            section.header.add_paragraph('Deep: {{deep_header_var}}')

    output_path = script_dir / 'template_many_headers.docx'
    doc.save(output_path)
    print(f"Created: {output_path}")


if __name__ == '__main__':
    print("Creating test fixture DOCX files...\n")
    create_sample_descr()
    create_simple_template()
    create_template_with_header()
    create_template_with_table()
    create_template_with_inflection()
    create_template_many_headers()
    print("\nAll fixtures created successfully!")
