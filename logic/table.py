# templates/logic/table.py
"""
Table handling user plugin.

Discovers tables in uploaded files, generates mapping UI,
and extracts tables as subdocs during rendering.

Convention-based dispatch:
  - match_table()  → discover tables in uploaded files
  - init_table()   → render mapping UI (preview + slot picker)
  - compute_table() → factory returning deferred extractors
"""
from docx import Document
from session_manager import get_file
from logic.subdoc_helper import extract_table_as_subdoc

# Which file-type variables contain tables to scan.
# Edit this list to add new sources: ['cv', 'portfolio'], etc.
TABLE_SOURCE_FILES = ['cv']


def match_table(session, lang_rules):
    """Discover tables in uploaded files listed in TABLE_SOURCE_FILES.

    Returns a list of structured discovery dicts for init_table() to render.
    Each dict carries a composite key (file access), preview data, and dimensions.
    init_table() receives this as `matched_list` — no file reading needed there.

    Composite key format: 'file_var:file_idx:table_idx' (e.g. 'cv:0:2').
    """
    results = []
    for file_var in TABLE_SOURCE_FILES:
        paths = get_file(session, file_var)
        if not paths:
            continue
        for file_idx, path in enumerate(paths):
            try:
                doc = Document(path)
                for t_idx, table in enumerate(doc.tables):
                    rows = len(table.rows)
                    cols = len(table.rows[0].cells) if rows > 0 else 0

                    # Preview: up to 3 rows of cell text
                    preview = []
                    for r in range(min(3, rows)):
                        preview.append([cell.text.strip() for cell in table.rows[r].cells])

                    results.append({
                        'key': f"{file_var}:{file_idx}:{t_idx}",
                        'preview': preview,
                        'rows': rows,
                        'cols': cols,
                    })
            except Exception:
                pass
    return results  # init_table() renders from this; empty list → no UI


def init_table(session, matched_list, variable_desc=None):
    """Generate table correspondence UI: previews + slot selection + hidden defaults.

    Iterates over DISCOVERED tables (from matched_list produced by match_table()),
    renders each with a preview and a <select> of available template slots
    (table_* vars from variable_desc). User picks which discovered table
    fills which slot — or skips it.

    matched_list: list of dicts from match_table():
      [{'key': 'cv:0:2', 'preview': [['Name','Age'],...], 'rows': 10, 'cols': 4}, ...]

    variable_desc: dict from descr.docx — used to discover table_* slot names.
      Not needed for data, only for slot labels in <select> options.

    Form fields:
      - <select name="table_map[N]"> → value is slot name ('table_1') or '' (skip)
      - <input name="table_key[N]" value="cv:0:2"> → composite key for discovered table
      - <input type="hidden" name="table_1" value=""> → ONE PER table_* SLOT from variable_desc
        Ensures all slots are in cleaned_data → excluded from missing_vars in computation.py.
        Count comes from descr.docx, NOT from discovered table count.
    """
    if not matched_list:
        return ''  # No tables discovered — nothing to render

    if not variable_desc:
        return '<div class="no-tables">variable_desc not available</div>'

    # Get table_* variables from descr.docx — these are template slots
    table_vars = sorted(
        [v for v in variable_desc if v.startswith('table_')],
        key=lambda v: int(v.split('_')[1])
    )
    if not table_vars:
        return '<div class="no-tables">No table slots defined in template</div>'

    # Hidden defaults for ALL table_* slots (from descr.docx, not discovered tables).
    # Ensures every slot is in cleaned_data → excluded from missing_vars.
    # compute_table() overwrites mapped slots; unmapped stay "".
    html = '<div class="table-mappings">'
    for tv in table_vars:
        html += f'<input type="hidden" name="{tv}" value="">'

    html += '<h3>Таблицы из файла → Маркеры в шаблоне:</h3>'
    html += '<p>Для каждой найденной таблицы выберите маркер в шаблоне:</p>'

    for idx, table_info in enumerate(matched_list):
        preview = table_info.get('preview', [])
        rows = table_info.get('rows', 0)
        cols = table_info.get('cols', 0)
        composite_key = table_info['key']  # 'cv:0:2'

        # Render preview table
        html += f'<div class="table-container">'
        html += f'<h4>Таблица #{idx+1} ({rows} строк, {cols} столбцов)</h4>'
        html += '<table>'
        for row in preview:
            html += '<tr>'
            for cell in row:
                html += f'<td>{cell}</td>'
            html += '</tr>'
        html += '</table>'

        # Slot pick emitted as a Jinja tag (not read in Python and baked in) so a
        # fresh pick doesn't move the emitted template / ui_hash. The framework
        # renders it post-hash (render_dynamic_ui_block); a structure-change
        # re-render refills it from the session (no value-loss).
        map_field = f"table_map[{idx}]"
        html += '<div><label>Маркер для этой таблицы:</label>'
        html += f'<select name="{map_field}" style="padding: 5px; margin-left: 5px;">'
        html += '<option value="">— Не использовать —</option>'

        for tv in table_vars:
            slot_label = variable_desc[tv].get('description', tv)
            sel = ("{% if session.get('" + map_field + "') == '" + tv + "' %}"
                   "selected{% endif %}")
            html += f'<option value="{tv}" {sel}>{slot_label}</option>'

        html += '</select></div>'

        # Hidden field to carry the composite key (file access info)
        html += f'<input type="hidden" name="table_key[{idx}]" value="{composite_key}">'

        html += '</div>'

    html += '</div>'
    return html


def compute_table(**kwargs):
    """Factory: returns {var_name: deferred_extractor} for all mapped tables.

    Session arrives via **kwargs pass-through from compute_variables().
    File paths resolved at compute time, captured in extractor closures.
    Composite key format: 'file_var:file_idx:table_idx' (e.g. 'cv:0:2').
    """
    session = kwargs.get('session')

    # Build mapping: slot_name → composite_key from table_map[] / table_key[]
    table_mappings = {}  # e.g. {'table_1': 'cv:0:2'}
    for key, value in kwargs.items():
        if key.startswith('table_map['):
            slot_name = value  # e.g. 'table_1'
            if not slot_name:
                continue  # User chose "Не использовать"
            idx = key.split('[')[1].rstrip(']')
            composite_key = kwargs.get(f'table_key[{idx}]', '')
            if composite_key:
                table_mappings[slot_name] = composite_key

    if not table_mappings:
        return {}

    def make_extractor(composite_key):
        # Parse 'file_var:file_idx:table_idx'
        file_var, file_idx, table_idx = composite_key.split(':')
        file_idx, table_idx = int(file_idx), int(table_idx)

        # Resolve file path at compute time — capture in closure
        uploaded_paths = get_file(session, file_var) if session else []
        file_path = uploaded_paths[file_idx] if uploaded_paths and file_idx < len(uploaded_paths) else None

        def _extract(tpl, **ctx):
            if not file_path:
                return tpl.new_subdoc()
            source = Document(file_path)
            return extract_table_as_subdoc(
                tpl, source, table_idx,
                context_data=ctx,
                render_placeholders=True,
            )
        return _extract

    return {
        slot_name: make_extractor(composite_key)
        for slot_name, composite_key in table_mappings.items()
    }
