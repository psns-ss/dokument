# logic/university.py
"""
University data processing logic for the document generator.
Covers matching, UI generation, and bookmark-based subdoc creation.
"""
from pathlib import Path
from typing import Dict, List, Any
from docx import Document
import logging
import re
import yaml
from logic.match_helpers import find_matches
from logic.subdoc_helper import create_subdoc_from_bookmark
from docxtpl import DocxTemplate, Subdoc

logger = logging.getLogger(__name__)


def _pack_file(name: str) -> Path:
    """A data file of the pack this plugin belongs to; root pack as fallback.

    A plugin always sits in `<pack>/logic/`, so two levels up is that pack's
    directory — no CWD-relative guessing, which silently read the ROOT pack's
    files from every pack. The fallback covers the host layout of the root pack:
    there `logic/` is at the repo root while the pack
    files live under `templates/`, whereas in Docker `logic/` is mounted inside
    `templates/` and the first candidate hits directly.
    """
    cand = Path(__file__).resolve().parent.parent / name
    return cand if cand.is_file() else Path("templates") / name

# A graduation year: a standalone 4-digit number in a plausible range
# (1950–2039). Used to pull the year out of a matched CV row.
_YEAR_RE = re.compile(r"\b(19[5-9]\d|20[0-3]\d)\b")


def _extract_year(cells: List[str]) -> str:
    """Return the first plausible graduation year in a row's cells, else ''."""
    for cell in cells:
        m = _YEAR_RE.search(cell)
        if m:
            return m.group(1)
    return ""


def match_university(
    session: Dict[str, Any],
    lang_rules: Dict[str, Any]
) -> List[str]:
    """
    Match university names from the uploaded CV files against the database.

    Reads CV file paths from session via get_file(session, 'cv').
    Returns empty list if no CV files are uploaded.

    Args:
        session: Flask session with uploaded_files dict.
        lang_rules: Language rules (not used here, passed for compatibility).

    Returns:
        List of keys of the matched universities.
    """
    from session_manager import get_file

    cv_paths = get_file(session, 'cv')
    if not cv_paths:
        logger.debug("No CV files uploaded — skipping university matching")
        return []

    cv_path = cv_paths[0]
    if not Path(cv_path).exists():
        logger.warning("CV file not found: %s", cv_path)
        return []

    doc = Document(cv_path)
    matched_keys = []
    cv_years: Dict[str, str] = {}

    # Load the university database
    try:
        db_path = _pack_file("university.yaml")
        with open(db_path, "r", encoding="utf-8") as f:
            universities_db = yaml.safe_load(f) or {}
    except Exception as e:
        logger.error("Error loading university.yaml: %s", e)
        return []

    # Extract text from all tables for matching
    all_text_from_tables = ""
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            row_text = " ".join(cells)
            all_text_from_tables += row_text + "\n"

            # Look for matches in the database
            found_keys = find_matches(
                raw_text=row_text,
                yaml_dict=universities_db,
                keys=("short", "full")
            )
            matched_keys.extend(found_keys)

            # The graduation year sits in the same row as the matched university
            # name — so only extract it for rows that actually matched (a work-
            # history row has years too, but it won't match a university).
            if found_keys:
                year = _extract_year(cells)
                if year:
                    for key in found_keys:
                        cv_years.setdefault(key, year)

    # Remove duplicates, preserving order
    seen = set()
    unique_keys = []
    for key in matched_keys:
        if key not in seen:
            seen.add(key)
            unique_keys.append(key)

    # Carry the extracted years to init_university (it reads session['cv_years']).
    # Merged, not replaced, so a re-run doesn't drop years for keys matched earlier.
    session.setdefault("cv_years", {}).update(cv_years)

    logger.debug("Matched universities: %s (years: %s)", unique_keys, cv_years)
    return unique_keys

def init_university(
    session: Dict[str, Any],
    matched_list: List[str],
    variable_desc: Dict[str, Dict] = None,
) -> str:
    """
    Generate an HTML form to select the matched universities and graduation years.

    Args:
        session: Flask session.
        matched_list: List of keys of the matched universities.

    Returns:
        HTML code to insert into the form.
    """
    try:
        db_path = _pack_file("university.yaml")
        with open(db_path, "r", encoding="utf-8") as f:
            universities_db = yaml.safe_load(f) or {}
    except Exception as e:
        logger.error("Error loading university.yaml: %s", e)
        return ""

    html_parts = [
        '<div class="matching-section" style="margin: 20px 0; padding: 15px; background-color: #f8f9fa; border: 1px solid #dee2e6; border-radius: 5px;">',
        '<h4>Ваш академический путь:</h4>'
    ]

    for i, key in enumerate(matched_list):
        if key not in universities_db:
            continue

        # The user's picks are emitted as Jinja placeholders (not read in Python
        # and baked in), so a fresh pick doesn't change the emitted template — and
        # thus ui_hash — killing the extra init-page click. The framework renders
        # them post-hash (render_dynamic_ui_block); a structure-change re-render
        # refills them from the session (no value-loss).
        order_field = f"university_order[{i}]"
        year_field = f"year_{i}"

        html_parts.append(
            f'<div style="margin-bottom: 10px; padding: 10px; border: 1px solid #ccc;">'
            f'<select name="{order_field}" style="padding: 5px; margin-right: 10px;">'
            f'<option value="">Выберите университет</option>'
        )

        for db_key, db_data in universities_db.items():
            if db_key not in matched_list:
                continue
            # stored pick (session) wins, else the matched default `key`.
            # Plain string -> the {% %} stays literal; Python values by concat.
            sel = ("{% if (session.get('" + order_field + "')"
                   " or '" + key + "') == '" + db_key + "' %}selected{% endif %}")
            html_parts.append(
                f'<option value="{db_key}" {sel}>{db_data.get("full", db_key)}</option>'
            )

        html_parts.append('</select>')
        # typed year wins; else CV year stashed during matching; else empty.
        # `| e` = html.escape, moved into the template.
        year = ("{{ (session.get('" + year_field + "')"
                " or session.get('cv_years', {}).get('" + key + "', '')) | e }}")
        html_parts.append(
            f'Год окончания: <input type="text" name="{year_field}" value="{year}" '
            f'style="padding: 5px; margin-right: 10px;">'
        )
        html_parts.append('</div>')

    html_parts.append('</div>')
    return "\n".join(html_parts)

def compute_university(**kwargs) -> List[Dict[str, Any]]:
    """
    Return a list of dicts with full data for the selected universities, including graduation years.

    Args:
        **kwargs: Form data, including university_order[*] and year_*.

    Returns:
        List of dicts with keys: key, short, full, historical, location, founded, specialties, year.
    """
    try:
        db_path = _pack_file("university.yaml")
        with open(db_path, "r", encoding="utf-8") as f:
            universities_db = yaml.safe_load(f) or {}
    except Exception as e:
        logger.error("Error loading university.yaml: %s", e)
        return []

    # Extract the university order from the form data
    order = []
    for key, value in kwargs.items():
        if key.startswith("university_order["):
            if value:
                order.append(value)

    # Build the list of dicts with full data
    result = []
    for i, key in enumerate(order):
        if key not in universities_db:
            continue

        uni_data = universities_db[key]
        uni_dict = {"key": key, **uni_data}

        # Add the graduation year from the form
        year_key = f"year_{i}"
        uni_dict["year"] = kwargs.get(year_key, "")  # Default value: empty string
#        if year_key in kwargs and kwargs[year_key]:
#            uni_dict["year"] = kwargs[year_key]

        result.append(uni_dict)

    logger.debug("Formed university data: %s", result)
    return result

def university_historical(
    tpl: DocxTemplate,
    **kwargs
) -> Dict[str, Subdoc]:
    """
    Generate subdocs for the selected universities based on bookmarks in university.docx.

    Args:
        tpl: DocxTemplate used to create the subdocs.
        **kwargs: Context with the university data (list of dicts).

    Returns:
        Dict {university_key: Subdoc} to insert into the template.
    """
    logger.debug("Generating subdocs for universities")
    db_path = _pack_file("university.yaml")
    src_path = _pack_file("university.docx")

    try:
        with open(db_path, "r", encoding="utf-8") as f:
            univ_db = yaml.safe_load(f) or {}
    except Exception as e:
        logger.error("Error loading university.yaml: %s", e)
        return {}

    # Extract the university list from the context
    university_list = [u["key"] for u in kwargs.get("university", [])]
    subdoc_dict = {}

    # Pass full kwargs — bookmark content determines which placeholders exist.
    # Extra keys in context are harmless; the hook doesn't need to pre-declare
    # which variables its DOCX uses.
    render_ctx = kwargs

    for key in university_list:
        if key not in univ_db:
            continue

        bookmark = univ_db[key].get("bookmark")
        if not bookmark:
            continue

        try:
            logger.debug("Creating subdoc for %s (bookmark: %s)", key, bookmark)
            subdoc_dict[key] = create_subdoc_from_bookmark(
                tpl, bookmark, src_path, render_context=render_ctx
            )
        except Exception as e:
            logger.error("Error creating subdoc for %s: %s", key, e)
            subdoc_dict[key] = tpl.new_subdoc()

    logger.debug("Subdocs created: %d", len(subdoc_dict))
    return subdoc_dict
